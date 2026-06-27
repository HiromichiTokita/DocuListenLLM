"""VOICEVOX テキスト読み上げ (AI Audio Director 対応版)"""
import json
import os
import queue
import re
import subprocess
import sys
import threading
import time
from tkinter import filedialog
from typing import Optional
import customtkinter as ctk
import numpy as np
import ollama
import pedalboard
import psutil
import pypdf
import requests
import sounddevice as sd
import soundfile as sf

import irodori_engine

try:
    from docx import Document as DocxDocument
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False


# ─────────────────────────────────────────
# 定数
# ─────────────────────────────────────────
VOICEVOX_PORT = 50021
VOICEVOX_URL  = f"http://127.0.0.1:{VOICEVOX_PORT}"
OLLAMA_URL   = "http://127.0.0.1:11434"

ENGINE_STARTUP_TIMEOUT = 60
TEMP_SLOT_COUNT        = 5
SILENCE_PADDING_SEC    = 0.4
PRE_PADDING_SEC        = 0.15

DEFAULT_SPEAKER_DATA: dict[str, dict[str, int]] = {
    "四国めたん": {"ノーマル": 2},
    "ずんだもん":  {"ノーマル": 3},
}
DEFAULT_NARRATOR_CHAR = "四国めたん"
DEFAULT_DIALOGUE_CHAR = "ずんだもん"

SPEED_MIN     = 0.5
SPEED_MAX     = 3.0
SPEED_DEFAULT = 1.0
APP_VERSION   = "v1.3.0"
DEBUG         = False  # True にするとデバッグ出力が有効になります
VOLUME_DEFAULT = 1.0

OLLAMA_MODEL_DEFAULT = "llama3.2"

DEFAULT_ARCHETYPES: dict[str, dict[str, str]] = {
    "ナレーション":       {"char": "四国めたん",          "style": "ノーマル"},
    "主人公 女":          {"char": "春日部つむぎ",         "style": "ノーマル"},
    "主人公 男":          {"char": "黒沢冴白",             "style": "ノーマル"},
    "子供 男":            {"char": "白上虎太郎",           "style": "ノーマル"},
    "子供 女":            {"char": "櫻歌ミコ",             "style": "ノーマル"},
    "若者 男":            {"char": "波音リツ",             "style": "ノーマル"},
    "若者 女":            {"char": "冥鳴ひまり",           "style": "ノーマル"},
    "中年 男":            {"char": "青山龍星",             "style": "ノーマル"},
    "中年 女":            {"char": "雨晴はう",             "style": "ノーマル"},
    "老人 男":            {"char": "ちび式じい",           "style": "ノーマル"},
    "老人 女":            {"char": "後鬼",                 "style": "ぬいぐるみ"},
    "ロボット":           {"char": "ナースロボ＿タイプＴ", "style": "ノーマル"},
    "人外仲間(かわいい)": {"char": "ずんだもん",           "style": "ノーマル"},
    "人外仲間(かっこいい)": {"char": "剣崎雌雄",          "style": "ノーマル"},
    "怪物":               {"char": "No.7",                 "style": "ノーマル"},
}

ARCHETYPES = list(DEFAULT_ARCHETYPES.keys())

QUOTE_OPTIONS = ["「」", "『』", "\"”", "()", "（）", "なし"]

PROFILING_PROMPT = """You are an Audio Director. Read the [Full Story Text].
Identify ALL the speaking characters and important named characters in the story. You must extract AS MANY characters as possible, do not stop at just one.
For each character, deduce their archetype from the exact 14 categories below:
["主人公 女", "主人公 男", "子供 男", "子供 女", "若者 男", "若者 女", "中年 男", "中年 女", "老人 男", "老人 女", "ロボット", "人外仲間(かわいい)", "人外仲間(かっこいい)", "怪物"]

You MUST output a JSON object with a single key "characters", whose value is a JSON array of objects. Each object must have exactly two keys: "name" and "category".
Example format:
{
  "characters": [
    {"name": "ハンネローレ", "category": "主人公 女"},
    {"name": "ラザンタルク", "category": "若者 男"},
    {"name": "コルドゥラ", "category": "中年 女"}
  ]
}
Do not output anything else.
"""

ATTRIBUTION_PROMPT = """You are an expert Audio Director.
Below is the Character Profile (dictionary) of this story:
{character_profile}

You will be given a [Broad Context] (sentences before and after the target dialogue) and the [Target Dialogue] itself.
Your task is to analyze the broad context (look both before and after the target dialogue for clues like "と〇〇が言った" or speaker conversational turns) to deduce WHO is speaking.

You MUST choose EXACTLY ONE category from this exact list (DO NOT use "ナレーション"):
["主人公 女", "主人公 男", "子供 男", "子供 女", "若者 男", "若者 女", "中年 男", "中年 女", "老人 男", "老人 女", "ロボット", "人外仲間(かわいい)", "人外仲間(かっこいい)", "怪物"]

Output ONLY a JSON object with a single key "category". Do not output anything else.
"""

STATUS_COLORS = {
    "idle":    ("gray50",     "gray60"),
    "working": ("blue3",      "dodger blue"),
    "ok":      ("green4",     "green3"),
    "error":   ("red3",       "red2"),
    "stopped": ("orange3",    "orange2"),
    "paused":  ("goldenrod3", "goldenrod1"),
}

COLOR_PAUSE_ACTIVE = ("#3B8ED0", "#1F6AA5")
COLOR_STOP_ACTIVE  = ("#C0392B", "#922B21")
COLOR_DISABLED     = ("gray75",  "gray35")

HIGHLIGHT_TAG = "chunk_active"
HIGHLIGHT_BG  = "#1F538A"
HIGHLIGHT_FG  = "#FFFFFF"

_PRIORITY_MAP: dict[str, int] = {
    "Normal":       psutil.NORMAL_PRIORITY_CLASS,
    "Below Normal": psutil.BELOW_NORMAL_PRIORITY_CLASS,
    "Low":          psutil.IDLE_PRIORITY_CLASS,
}


# ─────────────────────────────────────────
# パス解決
# ─────────────────────────────────────────
def _resolve_app_dir() -> str:
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def _resolve_engine_path() -> str:
    return os.path.join(_resolve_app_dir(), "engine", "run.exe")

_APP_DIR = _resolve_app_dir()

TEMP_PATHS: list[str] = [
    os.path.join(_APP_DIR, "temp_1.wav"),
    os.path.join(_APP_DIR, "temp_2.wav"),
    os.path.join(_APP_DIR, "temp_3.wav"),
    os.path.join(_APP_DIR, "temp_4.wav"),
    os.path.join(_APP_DIR, "temp_5.wav"),
]

SETTINGS_PATH = os.path.join(_APP_DIR, "settings.json")


# ─────────────────────────────────────────
# テキスト・音声ユーティリティ
# ─────────────────────────────────────────
def _speaker_for_chunk(chunk: str, narrator_id: int, dialogue_id: int,
                        valid_openers: tuple = ("「", "『")) -> int:
    return dialogue_id if chunk.startswith(valid_openers) else narrator_id

_NUM_KANJI = {
    "0": "ゼロ", "1": "いち", "2": "に", "3": "さん", "4": "よん",
    "5": "ご", "6": "ろく", "7": "なな", "8": "はち", "9": "きゅう",
}

def _digits_to_kana(text: str) -> str:
    """数字を1文字ずつかな読みに変換"""
    result = []
    for ch in text:
        result.append(_NUM_KANJI.get(ch, ch))
    return "".join(result)

def _rms_normalize(audio: np.ndarray, target_rms: float = 0.12) -> np.ndarray:
    rms = float(np.sqrt(np.mean(audio ** 2)))
    if rms < 1e-6:
        return audio
    return np.clip(audio * (target_rms / rms), -1.0, 1.0)

def _time_stretch(audio: np.ndarray, sample_rate: int, speed: float) -> np.ndarray:
    if abs(speed - 1.0) < 0.01:
        return audio
    was_1d = audio.ndim == 1
    buf = audio.reshape(1, -1).astype(np.float32) if was_1d else audio.T.astype(np.float32)
    stretched = pedalboard.time_stretch(buf, sample_rate, stretch_factor=speed)
    return stretched[0] if was_1d else stretched.T


# 再生はこのサンプルレートに統一する（VOICEVOXの24kHzは実績あり。
# Irodoriの48kHzは一部環境で sd.play が再生されないため、ここへリサンプルして同じ経路に載せる）
PLAYBACK_SR = 24000


def _resample_linear(audio: np.ndarray, sr_from: int, sr_to: int) -> np.ndarray:
    """numpy 線形補間によるリサンプル（依存追加なし）。speech 用途には十分。"""
    if sr_from == sr_to or sr_from <= 0 or sr_to <= 0 or audio.shape[0] == 0:
        return audio
    n_to = int(round(audio.shape[0] * sr_to / sr_from))
    if n_to <= 0:
        return audio
    x_old = np.linspace(0.0, 1.0, audio.shape[0], endpoint=False)
    x_new = np.linspace(0.0, 1.0, n_to, endpoint=False)
    if audio.ndim == 1:
        return np.interp(x_new, x_old, audio).astype(np.float32)
    chans = [np.interp(x_new, x_old, audio[:, c]) for c in range(audio.shape[1])]
    return np.stack(chans, axis=1).astype(np.float32)

def _best_style(styles: dict[str, int]) -> str:
    if "ノーマル" in styles:
        return "ノーマル"
    return min(styles, key=lambda name: styles[name])

def _convert_aozora_ruby(text: str, to_kana: bool) -> str:
    """青空文庫形式ルビ（親文字《よみ》, ｜親文字《よみ》）を処理する。

    to_kana=True なら親文字を読み（かな）に置換、False なら読みを除去して親文字を残す。
    """
    if to_kana:
        # ｜で範囲指定されたルビ: ｜親文字《よみ》 → よみ
        text = re.sub(r'｜[^《》]*《([^》]*)》', r'\1', text)
        # ｜無しのルビ: 直前の漢字列《よみ》 → よみ
        text = re.sub(r'[一-龥々〆ヵヶ]+《([^》]*)》', r'\1', text)
    # 上記で変換されなかった残りの《…》は除去し、ルビ開始記号｜も削除
    text = re.sub(r'《[^》]*》', '', text)
    text = re.sub(r'｜', '', text)
    return text


def _extract_text_from_file(filepath: str, ruby_to_kana: bool = False) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext in (".txt", ".md"):
        for encoding in ["utf-8", "shift_jis", "euc-jp"]:
            try:
                with open(filepath, "r", encoding=encoding, errors="replace") as f:
                    text = f.read()
                text = _convert_aozora_ruby(text, ruby_to_kana)
                # 目次行を除去（例: 「第一章　タイトル…………3」「…………」）
                text = re.sub(r'^[^\n]*[…・‥]{2,}\s*\d+\s*$', '', text, flags=re.MULTILINE)
                text = re.sub(r'^[…・‥\s]+$', '', text, flags=re.MULTILINE)
                text = re.sub(r'\n{3,}', '\n\n', text).strip()
                return text
            except UnicodeDecodeError:
                continue
        raise ValueError("テキストファイルのエンコーディングを判定できませんでした。")

    if ext == ".pdf":
        reader = pypdf.PdfReader(filepath)
        lines = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t:
                lines.append(t)
        return "\n".join(lines)

    if ext == ".docx":
        if not _DOCX_AVAILABLE:
            raise ImportError(
                "python-docx がインストールされていません。\n"
                "pip install python-docx を実行してください。"
            )
        doc = DocxDocument(filepath)
        lines = []
        for para in doc.paragraphs:
            try:
                for rt in para._element.xpath('.//w:rt'):
                    rt.getparent().remove(rt)
            except Exception:
                pass
            lines.append(para.text)
        return "\n".join(lines)

    if ext == ".epub":
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup
        except ImportError:
            raise ImportError(
                "EPUB読み込みモジュールがインストールされていません。\n"
                "pip install EbookLib beautifulsoup4 を実行してください。"
            )
        book = epub.read_epub(filepath)
        if DEBUG: print(f"[DEBUG] EPUB読み込み開始: {os.path.basename(filepath)}")

        # spine（読書順）に従ってドキュメントを並べる。spineに無いものは末尾へ補完。
        id_to_item = {it.get_id(): it for it in book.get_items()}
        ordered: list = []
        seen_ids: set = set()
        for entry in book.spine:
            idref = entry[0] if isinstance(entry, (list, tuple)) else entry
            it = id_to_item.get(idref)
            if it is not None and it.get_id() not in seen_ids:
                ordered.append(it)
                seen_ids.add(it.get_id())
        for it in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            if it.get_id() not in seen_ids:
                ordered.append(it)
                seen_ids.add(it.get_id())

        # 段落の区切りとして扱うブロック要素（インライン要素はここに含めない）
        _BLOCK_TAGS = [
            "p", "div", "li", "br", "blockquote", "section", "article",
            "h1", "h2", "h3", "h4", "h5", "h6", "tr", "figcaption", "caption", "pre",
        ]

        def _doc_text(item) -> str:
            """1つのドキュメントから本文テキストを頑健に抽出する。

            改行はブロック要素の境界にのみ挿入する。ルビ等のインライン要素
            （<ruby>漢<rt>かん</rt>…）内では区切らないため、ルビ除去後も
            「漢」「字」が1文字ずつ別段落に分かれてしまう問題を防ぐ。
            """
            raw = item.get_body_content() or b""
            if isinstance(raw, (bytes, bytearray)):
                raw = raw.decode("utf-8", errors="replace")
            soup = BeautifulSoup(raw, "html.parser")
            if ruby_to_kana:
                # ルビ付き部分のみ、親文字を読み(<rt>)に置換する。
                for ruby in soup.find_all("ruby"):
                    reading = "".join(rt.get_text() for rt in ruby.find_all("rt"))
                    if reading.strip():
                        ruby.replace_with(reading)
                    # 読みが空のルビはそのまま残し、下のrt除去で親文字を残す
            # ルビの読み・スクリプト等を除去（rb=ルビ親文字は残す）
            for tag in soup(["script", "style", "rt", "rp", "rtc"]):
                tag.decompose()
            # ブロック境界に改行マーカーを挿入してから区切り無しで連結する。
            # こうするとインラインのテキストノード（ルビ親文字含む）は連結され、
            # 改行は段落境界だけに入る。get_text は各ノードを1度だけ辿るので重複もしない。
            for tag in soup.find_all(_BLOCK_TAGS):
                if tag.name == "br":
                    tag.replace_with("\n")
                else:
                    tag.append("\n")
            txt = soup.get_text()
            lines = [ln.strip() for ln in txt.splitlines() if ln.strip()]
            return "\n".join(lines)

        chapters: list[str] = []
        nav_skipped = 0
        doc_count = sum(
            1 for it in ordered if it.get_type() == ebooklib.ITEM_DOCUMENT)
        for item in ordered:
            if item.get_type() != ebooklib.ITEM_DOCUMENT:
                continue
            # 自動生成のナビゲーション(目次)文書のみ除外（過剰スキップを避ける）
            if isinstance(item, epub.EpubNav):
                nav_skipped += 1
                if DEBUG: print(f"[DEBUG] EPUB: nav文書をスキップ '{item.get_name()}'")
                continue
            text = _doc_text(item)
            if text:
                chapters.append(text)
                if DEBUG: print(f"[DEBUG] EPUB: '{item.get_name()}' → {len(text)}文字")
            else:
                if DEBUG: print(f"[DEBUG] EPUB: '{item.get_name()}' → 空のためスキップ")

        result = "\n\n".join(chapters)
        if DEBUG: print(f"[DEBUG] EPUB: {len(chapters)}章 / 合計{len(result)}文字 "
              f"(ドキュメント数={doc_count}, nav除外={nav_skipped})")

        # 万一すべて空なら、何も除外せず全ドキュメントから再抽出（保険）
        if not result.strip():
            if DEBUG: print("[DEBUG] EPUB: 抽出結果が空 → フォールバックで全文を再抽出します。")
            fb = [_doc_text(it) for it in ordered
                  if it.get_type() == ebooklib.ITEM_DOCUMENT]
            result = "\n\n".join(t for t in fb if t)
            if DEBUG: print(f"[DEBUG] EPUB フォールバック結果: {len(result)}文字")

        # 青空文庫ルビ記法がテキストとして残った場合の保険
        result = _convert_aozora_ruby(result, ruby_to_kana)

        # 本文がほぼ取得できない場合の原因切り分け（画像のみEPUB等）
        if len(result.strip()) < 50:
            image_count = len(list(book.get_items_of_type(ebooklib.ITEM_IMAGE)))
            if DEBUG: print(f"[DEBUG] EPUB: 本文ほぼ無し (文字数={len(result.strip())}, "
                  f"画像数={image_count})")
            if image_count >= 5:
                raise ValueError(
                    f"このEPUBは画像形式（画像{image_count}枚）で本文の文字データが"
                    "無いため、読み上げできません。\n"
                    "テキスト形式のEPUBをご利用ください"
                    "（マンガ/スキャン等の画像本にはOCRが必要です）。")
            raise ValueError(
                "このEPUBから本文テキストを抽出できませんでした"
                "（DRM保護・非対応構造などの可能性があります）。")

        return result

    raise ValueError(f"非対応のファイル形式: {ext}")


# ─────────────────────────────────────────
# ログリダイレクター
# ─────────────────────────────────────────
import io

class _TeeStream(io.TextIOBase):
    """stdout/stderr を元のストリームとログテキストボックス両方に書き込む"""
    def __init__(self, original, textbox_getter):
        self._original = original
        self._get_tb = textbox_getter

    def write(self, s: str) -> int:
        if self._original is not None:
            try:
                self._original.write(s)
                self._original.flush()
            except Exception:
                pass
        tb = self._get_tb()
        if tb is not None:
            try:
                tb.configure(state="normal")
                tb.insert("end", s)
                tb.see("end")
                tb.configure(state="disabled")
            except Exception:
                pass
        return len(s)

    def flush(self):
        if self._original is not None:
            try:
                self._original.flush()
            except Exception:
                pass


# ─────────────────────────────────────────
# メインアプリケーションクラス
# ─────────────────────────────────────────
class VoicevoxTTSApp(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        self.title(f"VOICEVOX テキスト読み上げ (AI Audio Director) {APP_VERSION}")
        self.geometry("1000x900")
        self.minsize(800, 800)

        self._speaker_data: dict[str, dict[str, int]] = dict(DEFAULT_SPEAKER_DATA)

        self._stop_event  = threading.Event()
        self._pause_event = threading.Event()
        self._is_paused   = False
        self._is_playing  = False

        self._play_generation:  int = 0
        self._seek_pending_gen: int = 0

        self._audio_queue: Optional[queue.Queue] = None
        self._slot_semaphore = threading.Semaphore(4)

        self._producer_thread: Optional[threading.Thread] = None
        self._consumer_thread: Optional[threading.Thread] = None

        self._engine_proc: Optional[subprocess.Popen] = None

        self._chunks: list[str] = []
        self._chunk_tb_positions: list[tuple[str, str]] = []
        self._highlighted_chunk_idx: int = -1

        self._programmatic_slider_update = False

        self._http_session = requests.Session()
        self._script: list[dict] = []
        self._is_llm_processing = False
        self._char_dict_dirty = True
        self.char_rows: list[dict] = []

        # アーキタイプごとのキャラ/スタイル変数 (LLMモードで使用)
        self.archetype_vars: dict[str, dict] = {}
        self._archetype_menus: list[ctk.CTkOptionMenu] = []

        self._settings: dict = self._load_settings()
        # Irodori エンジン関連の既定値（後方互換: 既存設定を壊さず不足分のみ補完）
        _s = self._settings
        _s.setdefault("tts_engine", "voicevox")
        _s.setdefault("irodori_runtime_path", r"E:\project\IrodoriVDServer\irodori-vd-runtime")
        _s.setdefault("irodori_port", 8770)
        _s.setdefault("irodori_checkpoint", r"E:\project\DocuListenLLM\IrodoriTTS\model.safetensors")
        _s.setdefault("narrator_caption", irodori_engine.DEFAULT_NARRATOR_CAPTION)
        _caps = _s.setdefault("captions", {})
        for _cat, _cap in irodori_engine.DEFAULT_CAPTIONS.items():
            _caps.setdefault(_cat, _cap)
        _s.setdefault("irodori_use_ref", True)
        _s.setdefault("caption_seeds", {})
        _s.setdefault("category_engines", {})
        self._caption_seeds: dict[str, int] = dict(_s.get("caption_seeds", {}))
        self._category_engines: dict[str, str] = dict(_s.get("category_engines", {}))
        self._irodori = irodori_engine.IrodoriServerManager(
            runtime_path=_s.get("irodori_runtime_path", ""),
            port=int(_s.get("irodori_port", 8770)),
            checkpoint=_s.get("irodori_checkpoint", ""))
        self._recent_files: list[str] = self._settings.get("recent_files", [])
        _saved_port = self._settings.get("voicevox_port", VOICEVOX_PORT)
        global VOICEVOX_URL
        VOICEVOX_URL = f"http://127.0.0.1:{_saved_port}"
        self.quote1_var = ctk.StringVar(value=self._settings.get("quote1", "「」"))
        self.quote2_var = ctk.StringVar(value=self._settings.get("quote2", "『』"))
        self.quote3_var = ctk.StringVar(value=self._settings.get("quote3", "なし"))
        self.ruby_to_kana_var = ctk.BooleanVar(
            value=self._settings.get("ruby_to_kana", False))
        self._log_window: Optional[ctk.CTkToplevel] = None
        self._log_textbox: Optional[ctk.CTkTextbox] = None
        self._script_cache: dict[str, list] = {}
        self._last_positions: dict[str, int] = self._settings.get("last_positions", {})
        self._file_queue: list[str] = []

        import sys as _sys
        _sys.stdout = _TeeStream(_sys.stdout, lambda: self._log_textbox)
        _sys.stderr = _TeeStream(_sys.stderr, lambda: self._log_textbox)

        self._build_ui()
        self._set_controls_enabled(False)
        self._set_status("エンジン起動中...", "working")
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.bind("<Control-Up>",   self._on_volume_up)
        self.bind("<Control-Down>", self._on_volume_down)
        self.bind("<space>",        lambda e: self._on_space_key())
        self.bind("?",              lambda e: self._show_shortcuts())
        threading.Thread(target=self._start_engine_and_init, daemon=True).start()
        threading.Thread(target=self._fetch_ollama_models, daemon=True).start()

    # ─── 設定ファイル ─────────────────────────
    def _load_settings(self) -> dict:
        try:
            with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except (FileNotFoundError, json.JSONDecodeError):
            return {}

    def _save_settings(self) -> None:
        if not hasattr(self, "cpu_priority_var"):
            return
        archetype_settings: dict[str, dict] = {}
        for archetype, v in self.archetype_vars.items():
            archetype_settings[archetype] = {
                "char":  v["char"].get(),
                "style": v["style"].get(),
            }
        data = {
            "narrator_char":   self.narrator_char_var.get(),
            "narrator_style":  self.narrator_style_var.get(),
            "dialogue_char":   self.dialogue_char_var.get(),
            "dialogue_style":  self.dialogue_style_var.get(),
            "cpu_priority":    self.cpu_priority_var.get(),
            "playback_speed":  round(self.speed_slider.get(), 1),
            "silence_padding": round(self.padding_slider.get(), 1),
            "playback_volume":  round(self.volume_slider.get(), 2),
            "volume_boost":     self.volume_boost_var.get(),
            "normalize_volume": self.normalize_var.get() if hasattr(self, "normalize_var") else False,
            "ollama_model":    self.ollama_model_var.get(),
            "archetypes":      archetype_settings,
            "quote1":          self.quote1_var.get(),
            "quote2":          self.quote2_var.get(),
            "quote3":          self.quote3_var.get(),
            "ruby_to_kana":    self.ruby_to_kana_var.get(),
            "num_reading":     self.num_reading_var.get() if hasattr(self, "num_reading_var") else "そのまま",
            "custom_skip_patterns": getattr(self, "_custom_skip_patterns", []),
            "voicevox_port":   int(self.voicevox_port_var.get()) if hasattr(self, "voicevox_port_var") else VOICEVOX_PORT,
            "last_positions":  getattr(self, "_last_positions", {}),
            "recent_files":    getattr(self, "_recent_files", []),
            "character_rules": [
                {"name": r["name_var"].get().strip(), "category": r["cat_var"].get()}
                for r in getattr(self, "char_rows", [])
                if r["name_var"].get().strip()
            ],
            "tts_engine":           self.engine_var.get() if hasattr(self, "engine_var") else self._settings.get("tts_engine", "voicevox"),
            "irodori_runtime_path": self._settings.get("irodori_runtime_path", ""),
            "irodori_port":         self._settings.get("irodori_port", 8770),
            "irodori_checkpoint":   self._settings.get("irodori_checkpoint", ""),
            "narrator_caption":     self.narrator_caption_var.get() if hasattr(self, "narrator_caption_var") else self._settings.get("narrator_caption", ""),
            "captions":             {cat: var.get() for cat, var in self.caption_vars.items()} if hasattr(self, "caption_vars") else self._settings.get("captions", {}),
            "irodori_use_ref":      self.use_ref_var.get() if hasattr(self, "use_ref_var") else self._settings.get("irodori_use_ref", True),
            "caption_seeds":        {c: int(v) for c, v in self._caption_seeds.items()} if hasattr(self, "_caption_seeds") else self._settings.get("caption_seeds", {}),
            "category_engines":     dict(self._category_engines) if hasattr(self, "_category_engines") else self._settings.get("category_engines", {}),
        }
        try:
            with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except OSError:
            pass

    # ─── エンジンライフサイクル ───────────────
    def _start_engine_and_init(self) -> None:
        engine_exe = _resolve_engine_path()
        if os.path.isfile(engine_exe):
            try:
                self._engine_proc = subprocess.Popen(
                    [engine_exe],
                    creationflags=subprocess.CREATE_NO_WINDOW,
                )
            except OSError as exc:
                self.after(0, lambda e=str(exc): self._set_status(
                    f"エンジン起動失敗: {e}", "error"))
                self.after(0, lambda: self._set_controls_enabled(True))
                return
        else:
            self.after(0, lambda: self._set_status(
                "engine.exeが見つかりません。外部起動済みのVOICEVOXを確認中...", "working"))

        for elapsed in range(ENGINE_STARTUP_TIMEOUT):
            try:
                if self._http_session.get(
                        f"{VOICEVOX_URL}/version", timeout=2).status_code == 200:
                    break
            except Exception:
                pass

            if self._engine_proc and self._engine_proc.poll() is not None:
                self.after(0, lambda: self._set_status("エンジンが予期せず終了しました", "error"))
                self.after(0, lambda: self._set_controls_enabled(True))
                return

            self.after(0, lambda e=elapsed + 1: self._set_status(
                f"エンジン起動中... ({e}秒)", "working"))
            time.sleep(1)
        else:
            self.after(0, lambda: self._set_status(
                f"エンジンの起動が {ENGINE_STARTUP_TIMEOUT}秒 でタイムアウトしました", "error"))
            self.after(0, lambda: self._set_controls_enabled(True))
            return

        self.after(0, self._apply_engine_resource_limits)
        self._fetch_speakers()

    def _terminate_engine(self) -> None:
        if self._engine_proc and self._engine_proc.poll() is None:
            self._engine_proc.terminate()
        self._engine_proc = None

    def _on_close(self) -> None:
        try:
            self._save_settings()
        except Exception:
            pass
        self._terminate_engine()
        if hasattr(self, "_irodori") and self._irodori is not None:
            try:
                self._irodori.stop()
            except Exception:
                pass
        self.destroy()

    # ─── CPUリソース管理 ─────────────────────
    def _apply_engine_resource_limits(self) -> None:
        if self._engine_proc is None or self._engine_proc.poll() is not None:
            return
        try:
            proc = psutil.Process(self._engine_proc.pid)
            priority_str = self.cpu_priority_var.get()
            proc.nice(_PRIORITY_MAP.get(priority_str, psutil.BELOW_NORMAL_PRIORITY_CLASS))
        except (psutil.Error, AttributeError, ValueError, OSError):
            pass

    def _on_cpu_settings_changed(self, _: str = "") -> None:
        self._save_settings()
        self._apply_engine_resource_limits()

    # ─── Ollamaモデル一覧取得 ────────────────
    def _fetch_ollama_models(self) -> None:
        try:
            resp = requests.get(f"{OLLAMA_URL}/api/tags", timeout=5)
            resp.raise_for_status()
            models = [m["name"] for m in resp.json().get("models", [])]
            if not models:
                raise ValueError("empty")
        except Exception:
            models = ["llama3.2", "llama3"]
        self.after(0, lambda m=models: self._update_ollama_model_menu(m))

    def _update_ollama_model_menu(self, models: list[str]) -> None:
        saved = self._settings.get("ollama_model", OLLAMA_MODEL_DEFAULT)
        if saved not in models:
            models = [saved] + models if saved else models
        self.ollama_model_menu.configure(values=models)
        self.ollama_model_var.set(saved if saved in models else models[0])

    # ─── 話者一覧の動的取得 ──────────────────
    def _fetch_speakers(self) -> None:
        try:
            resp = self._http_session.get(f"{VOICEVOX_URL}/speakers", timeout=5)
            resp.raise_for_status()
            speaker_data: dict[str, dict[str, int]] = {}
            for character in resp.json():
                char_name = character["name"]
                styles: dict[str, int] = {
                    s["name"]: s["id"]
                    for s in character.get("styles", [])
                    if s.get("type", "talk") == "talk"
                }
                if styles:
                    speaker_data[char_name] = styles
            if speaker_data:
                self.after(0, lambda d=speaker_data: self._update_speaker_menus(d))
            else:
                self.after(0, lambda: self._set_status("待機中", "idle"))
                self.after(0, lambda: self._set_controls_enabled(True))
        except Exception:
            self.after(0, lambda: self._set_status("待機中", "idle"))
            self.after(0, lambda: self._set_controls_enabled(True))

    def _update_speaker_menus(self, speaker_data: dict[str, dict[str, int]]) -> None:
        self._speaker_data = speaker_data
        char_list = list(speaker_data.keys())

        def _resolve_char(saved: Optional[str], default: str) -> str:
            if saved and saved in speaker_data:
                return saved
            return default if default in speaker_data else char_list[0]

        narrator_char = _resolve_char(self._settings.get("narrator_char"), DEFAULT_NARRATOR_CHAR)
        dialogue_char = _resolve_char(self._settings.get("dialogue_char"), DEFAULT_DIALOGUE_CHAR)

        self.narrator_char_menu.configure(values=char_list)
        self.narrator_char_var.set(narrator_char)
        self.dialogue_char_menu.configure(values=char_list)
        self.dialogue_char_var.set(dialogue_char)

        self._apply_char_change(
            narrator_char, self.narrator_style_menu, self.narrator_style_var,
            saved_style=self._settings.get("narrator_style"),
        )
        self._apply_char_change(
            dialogue_char, self.dialogue_style_menu, self.dialogue_style_var,
            saved_style=self._settings.get("dialogue_style"),
        )

        self._update_archetype_menus(speaker_data)

        total_styles = sum(len(v) for v in speaker_data.values())
        self._set_status(
            f"話者一覧を取得しました ({len(char_list)}キャラ / {total_styles}スタイル)", "ok")
        self._set_controls_enabled(True)

    def _update_archetype_menus(self, speaker_data: dict[str, dict[str, int]]) -> None:
        char_list = list(speaker_data.keys())
        saved_archetypes: dict = self._settings.get("archetypes", {})
        for archetype, v in self.archetype_vars.items():
            saved = saved_archetypes.get(archetype, {})
            default_char = DEFAULT_ARCHETYPES.get(archetype, {}).get("char", DEFAULT_NARRATOR_CHAR)
            char = saved.get("char", default_char)
            if char not in speaker_data:
                char = char_list[0] if char_list else default_char
            v["char_menu"].configure(values=char_list)
            v["char"].set(char)
            self._apply_char_change(
                char, v["style_menu"], v["style"],
                saved_style=saved.get("style"),
            )

    def _apply_char_change(
        self,
        char: str,
        style_menu: ctk.CTkOptionMenu,
        style_var: ctk.StringVar,
        saved_style: Optional[str] = None,
    ) -> None:
        styles = self._speaker_data.get(char, {})
        if not styles:
            return
        style_menu.configure(values=list(styles.keys()))
        if saved_style and saved_style in styles:
            style_var.set(saved_style)
        else:
            style_var.set(_best_style(styles))

    def _on_narrator_char_changed(self, char: str) -> None:
        self._apply_char_change(char, self.narrator_style_menu, self.narrator_style_var)
        self._save_settings()

    def _on_dialogue_char_changed(self, char: str) -> None:
        self._apply_char_change(char, self.dialogue_style_menu, self.dialogue_style_var)
        self._save_settings()

    def _on_narrator_style_changed(self, _: str) -> None:
        self._save_settings()

    def _on_dialogue_style_changed(self, _: str) -> None:
        self._save_settings()

    def _get_speaker_id(self, char: str, style: str, fallback: int) -> int:
        return self._speaker_data.get(char, {}).get(style, fallback)

    def _get_narrator_id(self) -> int:
        return self._get_speaker_id(
            self.narrator_char_var.get(), self.narrator_style_var.get(), fallback=8)

    def _valid_openers(self) -> tuple:
        openers = []
        for var in [self.quote1_var, self.quote2_var, self.quote3_var]:
            val = var.get()
            if val and val != "なし" and len(val) == 2:
                openers.append(val[0])
        return tuple(openers) if openers else ("「", "『")

    def _is_dialogue(self, chunk: str) -> bool:
        return chunk.startswith(self._valid_openers())

    def _split_text(self, text: str) -> list[str]:
        if getattr(self, "num_reading_var", None) and self.num_reading_var.get() == "ひとつずつ":
            text = re.sub(r'\d+', lambda m: _digits_to_kana(m.group()), text)
        for pattern in getattr(self, "_custom_skip_patterns", []):
            try:
                text = re.sub(pattern, "", text)
            except re.error:
                pass
        valid_openers = self._valid_openers()
        quote_patterns = []
        for var in [self.quote1_var, self.quote2_var, self.quote3_var]:
            val = var.get()
            if val and val != "なし" and len(val) == 2:
                o, c = re.escape(val[0]), re.escape(val[1])
                quote_patterns.append(f"({o}[^{re.escape(val[1])}]*{c})")

        if not quote_patterns:
            sub = re.split(r"(?<=。)|\n", text)
            chunks = [c.strip() for c in sub if c.strip()]
            return self._merge_lone_brackets(chunks)

        pattern = "|".join(quote_patterns)
        parts = re.split(pattern, text)
        chunks: list[str] = []
        for part in parts:
            if not part or not part.strip():
                continue
            if part.startswith(valid_openers):
                chunks.append(part.strip())
            else:
                sub = re.split(r"(?<=。)|\n", part)
                chunks.extend(c.strip() for c in sub if c.strip())
        return self._merge_lone_brackets(chunks)

    def _merge_lone_brackets(self, chunks: list[str]) -> list[str]:
        LONE_OPEN  = {"（", "(", "『", "【", "〔"}
        LONE_CLOSE = {"）", ")", "』", "】", "〕"}
        LONE_SKIP  = {"…", "—", "―", "・", "、", "。", "・・・", "？", "！", "?", "!", "〜", "～"}
        result: list[str] = []
        i = 0
        while i < len(chunks):
            chunk = chunks[i]
            if chunk in LONE_SKIP:
                if result:
                    result[-1] += chunk
                i += 1
            elif chunk in LONE_OPEN:
                if i + 1 < len(chunks):
                    result.append(chunk + chunks[i + 1])
                    i += 2
                else:
                    i += 1  # 末尾の孤立開き括弧はスキップ
            elif chunk in LONE_CLOSE:
                if result:
                    result[-1] += chunk
                i += 1
            else:
                result.append(chunk)
                i += 1
        return result

    # ─── UI 構築 ─────────────────────────────
    def _build_ui(self) -> None:
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        self._build_top_bar()               # row 0
        self._build_text_frame()            # row 1
        self._build_navigation_frame()      # row 2
        self._build_control_frame()         # row 3
        self._build_history_bar()           # row 4
        self._build_casting_frame()         # row 5
        self._build_system_setting_frame()  # row 6

    def _build_top_bar(self) -> None:
        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 0))

        self.paste_play_btn = ctk.CTkButton(
            frame, text="▶ ペーストして再生", width=220,
            fg_color="#28a745", hover_color="#218838",
            command=self._on_paste_and_play)
        self.paste_play_btn.pack(side="left", padx=(0, 8))

        self.paste_btn = ctk.CTkButton(
            frame, text="📋 ペースト", width=120, command=self._on_paste_and_reset)
        self.paste_btn.pack(side="left", padx=(0, 14))

        self.import_btn = ctk.CTkButton(
            frame, text="📂 ファイル", width=120, command=self._on_import)
        self.import_btn.pack(side="left", padx=(0, 6))

        self.clear_btn = ctk.CTkButton(
            frame, text="🗑 クリア", width=100, command=self._on_clear)
        self.clear_btn.pack(side="left")

        self.ruby_kana_check = ctk.CTkCheckBox(
            frame, text="ルビ→かな", variable=self.ruby_to_kana_var,
            onvalue=True, offvalue=False)
        self.ruby_kana_check.pack(side="left", padx=(16, 0))

        ctk.CTkLabel(frame, text="エンジン:", width=60, anchor="e").pack(
            side="left", padx=(16, 2))
        self.engine_var = ctk.StringVar(value=self._settings.get("tts_engine", "voicevox"))
        ctk.CTkOptionMenu(
            frame, values=["voicevox", "irodori", "mixed"], variable=self.engine_var, width=110,
            command=lambda _v: self._on_engine_changed()).pack(side="left", padx=(0, 6))
        self.irodori_status_label = ctk.CTkLabel(
            frame, text="", font=ctk.CTkFont(size=11), text_color="gray60")
        self.irodori_status_label.pack(side="left")

        self.use_ref_var = ctk.BooleanVar(value=self._settings.get("irodori_use_ref", True))
        ctk.CTkCheckBox(frame, text="声を固定", variable=self.use_ref_var,
                        onvalue=True, offvalue=False,
                        command=self._save_settings).pack(side="left", padx=(8, 0))

        ctk.CTkButton(
            frame, text="🖥 ログ", width=80,
            fg_color="gray30", hover_color="gray20",
            command=self._open_log_window,
        ).pack(side="right")

        self._on_engine_changed()  # 初期状態ラベルを反映

    def _on_engine_changed(self) -> None:
        """エンジン切替時: 設定更新と状態ラベル表示（実際の起動は再生時）。"""
        eng = self.engine_var.get()
        self._settings["tts_engine"] = eng
        if eng == "irodori":
            self.irodori_status_label.configure(text="Irodori: 未起動")
        else:
            self.irodori_status_label.configure(text="")
        self._save_settings()

    def _reroll_voice(self, category: str) -> None:
        """カテゴリの声をリロール（seed変更）。次の再生で新しい基準音声になる。"""
        self._caption_seeds[category] = irodori_engine.new_seed()
        self._save_settings()
        label = "ナレーター" if category == "__narrator__" else category
        self._set_status(f"{label} の声をリロールしました（次の再生で反映）", "ok")

    def _set_category_engine(self, category: str, engine: str) -> None:
        """カテゴリ毎の使用エンジン（mixed時に有効）を更新・保存。"""
        self._category_engines[category] = engine
        self._save_settings()

    def _build_text_frame(self) -> None:
        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=(8, 0))
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(0, weight=1)

        self.textbox = ctk.CTkTextbox(frame, font=ctk.CTkFont(size=14), wrap="word")
        self.textbox.grid(row=0, column=0, sticky="nsew")
        self.textbox.insert("1.0", "ここにテキストを入力してください。")
        self._char_count_label = ctk.CTkLabel(
            frame, text="0 文字", anchor="e",
            font=ctk.CTkFont(size=11), text_color="gray60")
        self._char_count_label.grid(row=1, column=0, sticky="e", padx=4, pady=(1, 0))
        self.textbox._textbox.tag_configure(
            HIGHLIGHT_TAG, background=HIGHLIGHT_BG, foreground=HIGHLIGHT_FG)

        try:
            import windnd
            def _on_drop(files):
                if not files:
                    return
                f = files[0]
                filepath = f if isinstance(f, str) else f.decode(errors="replace")
                try:
                    text = _extract_text_from_file(
                        filepath, ruby_to_kana=self.ruby_to_kana_var.get())
                    self.textbox.delete("1.0", "end")
                    self.textbox.insert("1.0", text)
                    self._script = []
                    self._char_dict_dirty = True
                    self._clear_all_highlights()
                    self._update_chunks()
                    self.time_slider.set(0)
                    self._update_time_label(0)
                    self._highlight_chunk(0)
                    self._set_status(f"読み込み完了: {os.path.basename(filepath)}", "ok")
                    self._add_to_history(filepath)
                except Exception as e:
                    import traceback
                    if DEBUG: print("[DEBUG] !!! ドロップ読み込みエラー !!!")
                    traceback.print_exc()
                    self._set_status(f"ドロップ読み込みエラー: {e}", "error")
            windnd.hook_dropfiles(self.textbox._textbox, func=_on_drop, force_unicode=True)
        except ImportError:
            if DEBUG: print("[DEBUG] 'windnd' not installed. Drag-and-Drop disabled. Run: pip install windnd")

        self.textbox._textbox.bind("<KeyRelease>", self._on_text_changed)
        self.textbox._textbox.bind("<ButtonRelease-1>", self._on_textbox_click)

    def _build_navigation_frame(self) -> None:
        frame = ctk.CTkFrame(self, corner_radius=8)
        frame.grid(row=2, column=0, sticky="ew", padx=12, pady=(6, 0))
        frame.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(frame, text="位置:", width=45, anchor="e").grid(
            row=0, column=0, padx=(10, 4), pady=8, sticky="e")
        self.time_slider = ctk.CTkSlider(
            frame, from_=0, to=1, number_of_steps=1,
            command=self._on_time_slider_change)
        self.time_slider.set(0)
        self.time_slider.grid(row=0, column=1, padx=6, pady=8, sticky="ew")
        self.time_label = ctk.CTkLabel(frame, text="0 / 0", width=70, anchor="w")
        self.time_label.grid(row=0, column=2, padx=(4, 10), pady=8)

    def _build_casting_frame(self) -> None:
        self.tabview = ctk.CTkTabview(self, height=220)
        self.tabview.grid(row=5, column=0, sticky="ew", padx=12, pady=(6, 0))
        self.tabview.add("ルールベース")
        self.tabview.add("AIディレクター (15役)")

        self._build_rule_based_tab(self.tabview.tab("ルールベース"))
        self._build_llm_archetypes_tab(self.tabview.tab("AIディレクター (15役)"))

    def _build_rule_based_tab(self, parent: ctk.CTkFrame) -> None:
        char_list = list(self._speaker_data.keys())
        narrator_styles = list(
            self._speaker_data.get(DEFAULT_NARRATOR_CHAR, {"ノーマル": 2}).keys())
        dialogue_styles = list(
            self._speaker_data.get(DEFAULT_DIALOGUE_CHAR, {"ノーマル": 3}).keys())

        init_narrator_char  = self._settings.get("narrator_char",  DEFAULT_NARRATOR_CHAR)
        init_narrator_style = self._settings.get(
            "narrator_style",
            _best_style(self._speaker_data.get(DEFAULT_NARRATOR_CHAR, {"ノーマル": 2})))
        init_dialogue_char  = self._settings.get("dialogue_char",  DEFAULT_DIALOGUE_CHAR)
        init_dialogue_style = self._settings.get(
            "dialogue_style",
            _best_style(self._speaker_data.get(DEFAULT_DIALOGUE_CHAR, {"ノーマル": 3})))

        inner = ctk.CTkFrame(parent, fg_color="transparent")
        inner.pack(padx=6, pady=(6, 6), anchor="w")

        ctk.CTkLabel(inner, text="地の文:", width=55, anchor="e").grid(
            row=0, column=0, padx=(0, 4), pady=3, sticky="e")
        self.narrator_char_var = ctk.StringVar(value=init_narrator_char)
        self.narrator_char_menu = ctk.CTkOptionMenu(
            inner, values=char_list, variable=self.narrator_char_var,
            width=180, command=self._on_narrator_char_changed)
        self.narrator_char_menu.grid(row=0, column=1, padx=(0, 6), pady=3)

        ctk.CTkLabel(inner, text="スタイル:", width=60, anchor="e").grid(
            row=0, column=2, padx=(0, 4), pady=3, sticky="e")
        self.narrator_style_var = ctk.StringVar(value=init_narrator_style)
        self.narrator_style_menu = ctk.CTkOptionMenu(
            inner, values=narrator_styles, variable=self.narrator_style_var,
            width=160, command=self._on_narrator_style_changed)
        self.narrator_style_menu.grid(row=0, column=3, pady=3)

        ctk.CTkLabel(inner, text="セリフ:", width=55, anchor="e").grid(
            row=1, column=0, padx=(0, 4), pady=3, sticky="e")
        self.dialogue_char_var = ctk.StringVar(value=init_dialogue_char)
        self.dialogue_char_menu = ctk.CTkOptionMenu(
            inner, values=char_list, variable=self.dialogue_char_var,
            width=180, command=self._on_dialogue_char_changed)
        self.dialogue_char_menu.grid(row=1, column=1, padx=(0, 6), pady=3)

        ctk.CTkLabel(inner, text="スタイル:", width=60, anchor="e").grid(
            row=1, column=2, padx=(0, 4), pady=3, sticky="e")
        self.dialogue_style_var = ctk.StringVar(value=init_dialogue_style)
        self.dialogue_style_menu = ctk.CTkOptionMenu(
            inner, values=dialogue_styles, variable=self.dialogue_style_var,
            width=160, command=self._on_dialogue_style_changed)
        self.dialogue_style_menu.grid(row=1, column=3, pady=3)

    def _build_llm_archetypes_tab(self, parent: ctk.CTkFrame) -> None:
        inner_tabs = ctk.CTkTabview(parent)
        inner_tabs.pack(fill="both", expand=True, padx=4, pady=4)
        inner_tabs.add("配役リスト")
        inner_tabs.add("キャラクター辞書")

        # ── 配役リスト タブ ──
        cast_tab = inner_tabs.tab("配役リスト")
        scroll = ctk.CTkScrollableFrame(cast_tab, height=180)
        scroll.pack(fill="both", expand=True, padx=4, pady=4)

        char_list = list(self._speaker_data.keys())
        saved_archetypes: dict = self._settings.get("archetypes", {})
        self._archetype_menus = []
        self.caption_vars: dict[str, ctk.StringVar] = {}

        for row_idx, archetype in enumerate(ARCHETYPES):
            saved = saved_archetypes.get(archetype, {})
            archetype_default = DEFAULT_ARCHETYPES.get(archetype, {})
            init_char  = saved.get("char",  archetype_default.get("char",  DEFAULT_NARRATOR_CHAR))
            init_style = saved.get("style", archetype_default.get("style", "ノーマル"))

            char_var  = ctk.StringVar(value=init_char)
            style_var = ctk.StringVar(value=init_style)

            ctk.CTkLabel(scroll, text=f"{archetype}:", width=80, anchor="e").grid(
                row=row_idx, column=0, padx=(0, 4), pady=2, sticky="e")

            char_menu = ctk.CTkOptionMenu(
                scroll, values=char_list, variable=char_var, width=170,
                command=lambda c, sv=style_var, sm_ref=[None]: (
                    self._apply_char_change(c, sm_ref[0], sv) if sm_ref[0] else None
                ))
            char_menu.grid(row=row_idx, column=1, padx=(0, 6), pady=2)

            style_styles = list(
                self._speaker_data.get(init_char, {"ノーマル": 2}).keys())
            style_menu = ctk.CTkOptionMenu(
                scroll, values=style_styles, variable=style_var, width=130,
                command=lambda _: self._save_settings())
            style_menu.grid(row=row_idx, column=2, pady=2)

            char_menu.configure(
                command=lambda c, sv=style_var, sm=style_menu: (
                    self._apply_char_change(c, sm, sv),
                    self._save_settings(),
                ))

            # Irodori 用キャプション欄（VOICEVOX話者選択と共存）。
            # 「ナレーション」は下部の専用ナレーター欄に一本化するためスキップ。
            if archetype in irodori_engine.DEFAULT_CAPTIONS:
                cap_init = self._settings.get("captions", {}).get(
                    archetype, irodori_engine.DEFAULT_CAPTIONS.get(archetype, ""))
                cap_var = ctk.StringVar(value=cap_init)
                ctk.CTkEntry(scroll, textvariable=cap_var, width=320,
                             placeholder_text="Irodori キャプション").grid(
                    row=row_idx, column=3, padx=(8, 0), pady=2, sticky="we")
                self.caption_vars[archetype] = cap_var
                ctk.CTkButton(scroll, text="🎲", width=32,
                              command=lambda c=archetype: self._reroll_voice(c)).grid(
                    row=row_idx, column=4, padx=(4, 0), pady=2)
                _eng_var = ctk.StringVar(
                    value=self._category_engines.get(archetype, "voicevox"))
                ctk.CTkOptionMenu(
                    scroll, values=["voicevox", "irodori"], variable=_eng_var, width=92,
                    command=lambda v, c=archetype: self._set_category_engine(c, v)).grid(
                    row=row_idx, column=5, padx=(4, 0), pady=2)

            self.archetype_vars[archetype] = {
                "char":       char_var,
                "style":      style_var,
                "char_menu":  char_menu,
                "style_menu": style_menu,
            }
            self._archetype_menus.extend([char_menu, style_menu])

        # ナレーター（地の文）用 Irodori キャプション
        _nar_row = len(ARCHETYPES)
        ctk.CTkLabel(scroll, text="ナレーター:", width=80, anchor="e").grid(
            row=_nar_row, column=0, padx=(0, 4), pady=2, sticky="e")
        self.narrator_caption_var = ctk.StringVar(
            value=self._settings.get("narrator_caption", irodori_engine.DEFAULT_NARRATOR_CAPTION))
        ctk.CTkEntry(scroll, textvariable=self.narrator_caption_var, width=320,
                     placeholder_text="ナレーター キャプション").grid(
            row=_nar_row, column=1, columnspan=3, padx=(0, 0), pady=2, sticky="we")
        ctk.CTkButton(scroll, text="🎲", width=32,
                      command=lambda: self._reroll_voice("__narrator__")).grid(
            row=_nar_row, column=4, padx=(4, 0), pady=2)
        _nar_eng_var = ctk.StringVar(
            value=self._category_engines.get("__narrator__", "voicevox"))
        ctk.CTkOptionMenu(
            scroll, values=["voicevox", "irodori"], variable=_nar_eng_var, width=92,
            command=lambda v: self._set_category_engine("__narrator__", v)).grid(
            row=_nar_row, column=5, padx=(4, 0), pady=2)

        # ── キャラクター辞書 タブ ──
        dict_tab = inner_tabs.tab("キャラクター辞書")

        dict_header = ctk.CTkFrame(dict_tab, fg_color="transparent")
        dict_header.pack(fill="x", padx=10, pady=(10, 4))

        ctk.CTkLabel(dict_header, text="キャラクター辞書 (固定配役):",
                     font=ctk.CTkFont(size=12, weight="bold")).pack(side="left")

        self.extract_btn = ctk.CTkButton(
            dict_header, text="AIで自動抽出", width=110,
            command=self._start_extraction)
        self.extract_btn.pack(side="right", padx=(5, 0))

        ctk.CTkButton(
            dict_header, text="＋ 追加", width=65,
            command=lambda: self._add_char_row()).pack(side="right")

        header_frame = ctk.CTkFrame(dict_tab, fg_color="transparent")
        header_frame.pack(fill="x", padx=15, pady=(5, 0))
        ctk.CTkLabel(header_frame, text="キャラ名 (編集可)",
                     font=ctk.CTkFont(size=12, weight="bold"), width=160, anchor="w").pack(side="left", padx=5)
        ctk.CTkLabel(header_frame, text="配役カテゴリ",
                     font=ctk.CTkFont(size=12, weight="bold"), width=140, anchor="w").pack(side="left", padx=5)

        self.char_list_frame = ctk.CTkScrollableFrame(
            dict_tab, height=200,
            fg_color=("gray95", "gray15"),
            border_width=1, border_color="gray70",
        )
        self.char_list_frame.pack(fill="x", padx=10, pady=(0, 10))

        for rule in self._settings.get("character_rules", []):
            if isinstance(rule, dict):
                self._add_char_row(rule.get("name", ""), rule.get("category", "主人公 女"))

    # ─── キャラクター辞書 ─────────────────────
    _CAT_OPTIONS = [
        "主人公 女", "主人公 男", "子供 男", "子供 女",
        "若者 男", "若者 女", "中年 男", "中年 女",
        "老人 男", "老人 女", "ロボット",
        "人外仲間(かわいい)", "人外仲間(かっこいい)", "怪物",
    ]

    def _add_char_row(self, name: str = "", category: str = "主人公 女") -> None:
        if category not in self._CAT_OPTIONS:
            category = "主人公 女"
        row = ctk.CTkFrame(self.char_list_frame, fg_color="transparent")
        row.pack(fill="x", pady=3)

        name_var = ctk.StringVar(value=name)
        cat_var  = ctk.StringVar(value=category)

        entry = ctk.CTkEntry(row, textvariable=name_var, width=160,
                             placeholder_text="キャラ名 (例: 魔王)")
        entry.pack(side="left", padx=5)
        entry.bind("<KeyRelease>", lambda e: self._save_settings())

        opt = ctk.CTkOptionMenu(row, values=self._CAT_OPTIONS, variable=cat_var,
                                width=140, command=lambda _: self._save_settings())
        opt.pack(side="left", padx=5)

        row_data = {"frame": row, "name_var": name_var, "cat_var": cat_var}
        del_btn = ctk.CTkButton(
            row, text="削除", width=50,
            fg_color="#ff5555", hover_color="#cc4444",
            command=lambda r=row_data: self._delete_char_row(r))
        del_btn.pack(side="left", padx=(15, 0))

        self.char_rows.append(row_data)
        self._save_settings()

    def _delete_char_row(self, row_data: dict) -> None:
        row_data["frame"].destroy()
        self.char_rows = [r for r in self.char_rows if r is not row_data]
        self._save_settings()

    def _get_character_profile_json(self) -> str:
        rules = [
            {"name": r["name_var"].get().strip(), "category": r["cat_var"].get()}
            for r in self.char_rows
            if r["name_var"].get().strip()
        ]
        return json.dumps(rules, ensure_ascii=False)

    def _start_extraction(self) -> None:
        if self._is_llm_processing:
            self._set_status("AIの終了処理を待っています。数秒お待ちください", "working")
            return
        self._is_llm_processing = True
        self.extract_btn.configure(state="disabled", text="抽出中...")
        self._set_status("AI: 登場人物をプロファイリング中...", "working")
        threading.Thread(target=self._extraction_thread, daemon=True).start()

    def _clear_all_rows(self) -> None:
        for r in list(self.char_rows):
            r["frame"].destroy()
        self.char_rows.clear()

    def _extraction_thread(self) -> None:
        try:
            model     = self.ollama_model_var.get().strip() or OLLAMA_MODEL_DEFAULT
            full_text = self.textbox.get("1.0", "end").strip()
            if not full_text:
                return

            # コンテキスト溢れ防止: 序盤8000文字に制限
            max_chars = 8000
            if len(full_text) > max_chars:
                full_text = full_text[:max_chars] + "\n...[Text truncated to prevent overflow]..."
                if DEBUG: print(f"[DEBUG] Text truncated to {max_chars} chars for extraction.")

            import concurrent.futures as _cf
            _chat_args = dict(
                model=model,
                messages=[
                    {"role": "system", "content": PROFILING_PROMPT},
                    {"role": "user",   "content": f"[Full Story Text]\n{full_text}"},
                ],
                format="json",
                options={"num_ctx": 16384},
            )
            with _cf.ThreadPoolExecutor(max_workers=1) as _pool:
                _fut = _pool.submit(ollama.chat, **_chat_args)
                try:
                    resp = _fut.result(timeout=120)
                except _cf.TimeoutError:
                    raise TimeoutError("LLM抽出が120秒でタイムアウトしました")
            content = (resp.message.content
                       if hasattr(resp, "message")
                       else resp["message"]["content"])
            content = content.strip()
            if DEBUG: print(f"[DEBUG] Raw LLM Extraction Output:\n{content}\n" + "-" * 30)

            if not content:
                raise ValueError("LLM returned an empty response. (Context window might be too small)")

            if "```json" in content:
                content = content.split("```json")[1].split("```")[0].strip()
            elif "```" in content:
                content = content.split("```")[1].split("```")[0].strip()
            parsed = json.loads(content)

            # "characters" キーを優先して取り出す
            if isinstance(parsed, dict):
                if "characters" in parsed and isinstance(parsed["characters"], list):
                    extracted_list = parsed["characters"]
                else:
                    extracted_list = next(
                        (v for v in parsed.values() if isinstance(v, list)), []
                    )
            elif isinstance(parsed, list):
                extracted_list = parsed
            else:
                extracted_list = []

            if DEBUG: print(f"[DEBUG] Extraction parsed successfully. {len(extracted_list)} characters found.")
            self.after(0, self._clear_all_rows)
            for item in extracted_list:
                if isinstance(item, dict) and "name" in item and "category" in item:
                    self.after(0, lambda n=str(item["name"]), c=str(item["category"]):
                               self._add_char_row(n, c))
            self.after(0, lambda count=len(extracted_list):
                       self._set_status(f"抽出完了: {count}人のキャラをリストに追加しました", "ok"))

        except Exception as e:
            if DEBUG: print(f"[DEBUG] Extraction FAILED: {e}")
            self.after(0, lambda err=str(e): self._set_status(f"抽出失敗: {err}", "stopped"))
        finally:
            self._is_llm_processing = False
            self.after(0, lambda: self.extract_btn.configure(state="normal", text="AIで自動抽出"))

    def _build_system_setting_frame(self) -> None:
        frame = ctk.CTkFrame(self, corner_radius=8)
        frame.grid(row=6, column=0, sticky="ew", padx=12, pady=(6, 0))

        init_priority = self._settings.get("cpu_priority", "Below Normal")
        init_model    = self._settings.get("ollama_model", OLLAMA_MODEL_DEFAULT)
        init_port     = str(self._settings.get("voicevox_port", VOICEVOX_PORT))

        inner = ctk.CTkFrame(frame, fg_color="transparent")
        inner.pack(padx=10, pady=(8, 8), anchor="w")

        ctk.CTkLabel(inner, text="CPU優先度:", width=80, anchor="e").grid(
            row=0, column=0, padx=(0, 4), pady=3, sticky="e")
        self.cpu_priority_var = ctk.StringVar(value=init_priority)
        ctk.CTkOptionMenu(
            inner,
            values=["Normal", "Below Normal", "Low"],
            variable=self.cpu_priority_var,
            width=140,
            command=self._on_cpu_settings_changed,
        ).grid(row=0, column=1, pady=3, padx=(0, 20))

        ctk.CTkLabel(inner, text="Ollamaモデル:", width=95, anchor="e").grid(
            row=0, column=2, padx=(0, 4), pady=3, sticky="e")
        self.ollama_model_var = ctk.StringVar(value=init_model)
        self.ollama_model_menu = ctk.CTkOptionMenu(
            inner,
            values=[init_model],
            variable=self.ollama_model_var,
            width=180,
            command=lambda _: self._save_settings(),
        )
        self.ollama_model_menu.grid(row=0, column=3, pady=3)
        ctk.CTkButton(inner, text="速度テスト", width=80,
                      command=self._benchmark_ollama).grid(row=0, column=4, padx=(6, 0), pady=3)

        ctk.CTkLabel(inner, text="VOICEVOXポート:", width=110, anchor="e").grid(
            row=1, column=0, padx=(0, 4), pady=3, sticky="e")
        self.voicevox_port_var = ctk.StringVar(value=init_port)
        port_entry = ctk.CTkEntry(inner, textvariable=self.voicevox_port_var, width=80)
        port_entry.grid(row=1, column=1, pady=3, padx=(0, 20))
        port_entry.bind("<FocusOut>", lambda _: self._apply_voicevox_port())
        port_entry.bind("<Return>",   lambda _: self._apply_voicevox_port())

        ctk.CTkLabel(inner, text="数字の読み方:", width=110, anchor="e").grid(
            row=1, column=2, padx=(0, 4), pady=3, sticky="e")
        self.num_reading_var = ctk.StringVar(
            value=self._settings.get("num_reading", "そのまま"))
        ctk.CTkOptionMenu(
            inner, values=["そのまま", "ひとつずつ"],
            variable=self.num_reading_var, width=120,
            command=lambda _: self._save_settings(),
        ).grid(row=1, column=3, pady=3)

        # カスタム除去パターン
        self._custom_skip_patterns: list[str] = self._settings.get("custom_skip_patterns", [])
        skip_frame = ctk.CTkFrame(frame, fg_color="transparent")
        skip_frame.pack(padx=10, pady=(0, 4), fill="x")
        ctk.CTkLabel(skip_frame, text="除去パターン(正規表現):", width=150, anchor="w",
                     font=ctk.CTkFont(size=11)).pack(side="left", padx=(0, 6))
        self._skip_pattern_entry = ctk.CTkEntry(skip_frame, width=240,
            placeholder_text="例: \\[.*?\\]")
        self._skip_pattern_entry.pack(side="left", padx=(0, 6))
        ctk.CTkButton(skip_frame, text="追加", width=50,
                      command=self._add_skip_pattern).pack(side="left", padx=(0, 4))
        ctk.CTkButton(skip_frame, text="クリア", width=60,
                      command=self._clear_skip_patterns).pack(side="left")
        self._skip_pattern_label = ctk.CTkLabel(skip_frame, text=self._skip_patterns_text(),
            font=ctk.CTkFont(size=10), text_color="gray60")
        self._skip_pattern_label.pack(side="left", padx=(8, 0))

        quote_frame = ctk.CTkFrame(frame, fg_color="transparent")
        quote_frame.pack(padx=10, pady=(0, 8), fill="x")

        ctk.CTkLabel(quote_frame, text="会話括弧1:", width=65, anchor="e").pack(side="left", padx=(0, 4))
        ctk.CTkOptionMenu(quote_frame, values=QUOTE_OPTIONS, variable=self.quote1_var,
                          width=80, command=lambda _: self._save_settings()).pack(side="left")

        ctk.CTkLabel(quote_frame, text="会話括弧2:", width=65, anchor="e").pack(side="left", padx=(10, 4))
        ctk.CTkOptionMenu(quote_frame, values=QUOTE_OPTIONS, variable=self.quote2_var,
                          width=80, command=lambda _: self._save_settings()).pack(side="left")

        ctk.CTkLabel(quote_frame, text="会話括弧3:", width=65, anchor="e").pack(side="left", padx=(10, 4))
        ctk.CTkOptionMenu(quote_frame, values=QUOTE_OPTIONS, variable=self.quote3_var,
                          width=80, command=lambda _: self._save_settings()).pack(side="left")

    def _text_hash(self) -> str:
        import hashlib
        text = self.textbox.get("1.0", "end-1c")
        return hashlib.md5(text.encode("utf-8", errors="replace")).hexdigest()[:12]

    def _save_last_position(self) -> None:
        h = self._text_hash()
        idx = int(self.time_slider.get())
        self._last_positions[h] = idx
        # 古いエントリを削除して100件以内に保つ
        if len(self._last_positions) > 100:
            oldest = next(iter(self._last_positions))
            del self._last_positions[oldest]
        self._save_settings()

    def _restore_last_position(self) -> None:
        h = self._text_hash()
        idx = self._last_positions.get(h, 0)
        if idx > 0 and self._chunks:
            idx = min(idx, len(self._chunks) - 1)
            self.time_slider.set(idx)
            self._update_time_label(idx)
            self._clear_all_highlights()
            self._highlight_chunk(idx)

    def _show_queue_window(self) -> None:
        win = ctk.CTkToplevel(self)
        win.title("連続再生キュー")
        win.geometry("520x360")
        win.grid_columnconfigure(0, weight=1)
        win.grid_rowconfigure(1, weight=1)

        btn_frame = ctk.CTkFrame(win, fg_color="transparent")
        btn_frame.grid(row=0, column=0, sticky="ew", padx=8, pady=(8, 4))
        ctk.CTkButton(btn_frame, text="＋ ファイル追加", width=110,
                      command=lambda: self._queue_add_files(scroll_frame, win)).pack(side="left", padx=(0, 6))
        ctk.CTkButton(btn_frame, text="クリア", width=70,
                      command=lambda: self._queue_clear(scroll_frame)).pack(side="left", padx=(0, 6))
        ctk.CTkButton(btn_frame, text="▶ キュー再生", width=110,
                      command=lambda: (win.destroy(), self._start_queue_playback())).pack(side="left")

        scroll_frame = ctk.CTkScrollableFrame(win, height=260)
        scroll_frame.grid(row=1, column=0, sticky="nsew", padx=8, pady=(0, 8))
        scroll_frame.grid_columnconfigure(0, weight=1)
        self._refresh_queue_display(scroll_frame)
        win.grab_set()

    def _refresh_queue_display(self, frame: ctk.CTkScrollableFrame) -> None:
        for w in frame.winfo_children():
            w.destroy()
        for i, path in enumerate(self._file_queue):
            ctk.CTkLabel(frame, text=f"{i+1}. {os.path.basename(path)}", anchor="w").grid(
                row=i, column=0, sticky="w", padx=4, pady=1)

    def _queue_add_files(self, frame: ctk.CTkScrollableFrame, win: ctk.CTkToplevel) -> None:
        paths = filedialog.askopenfilenames(
            filetypes=[("Text/EPUB", "*.txt *.epub *.md"), ("All", "*.*")],
            title="ファイルをキューに追加",
        )
        for p in paths:
            if p not in self._file_queue:
                self._file_queue.append(p)
        self._refresh_queue_display(frame)

    def _queue_clear(self, frame: ctk.CTkScrollableFrame) -> None:
        self._file_queue.clear()
        self._refresh_queue_display(frame)

    def _start_queue_playback(self) -> None:
        if not self._file_queue:
            self._set_status("キューが空です", "error")
            return
        threading.Thread(target=self._queue_playback_thread, daemon=True).start()

    def _queue_playback_thread(self) -> None:
        for path in list(self._file_queue):
            if not os.path.exists(path):
                continue
            try:
                text = _extract_text_from_file(path, ruby_to_kana=self.ruby_to_kana_var.get())
            except Exception as exc:
                self.after(0, lambda: self._set_status(f"キュー読み込みエラー: {exc}", "error"))
                continue
            self.after(0, lambda t=text, p=path: self._load_text_for_queue(t, p))
            # 再生終了まで待機
            import time as _t
            _t.sleep(0.5)
            while self._is_playing or self._is_llm_processing:
                _t.sleep(0.3)
            if self._stop_event.is_set():
                break
        self.after(0, lambda: self._set_status("キュー再生完了", "ok"))

    def _load_text_for_queue(self, text: str, path: str) -> None:
        self.textbox.delete("1.0", "end")
        self.textbox.insert("1.0", text)
        self._script = []
        self._script_cache.clear()
        self._char_dict_dirty = True
        self._clear_all_highlights()
        self._update_chunks()
        self.time_slider.set(0)
        self._update_time_label(0)
        self._highlight_chunk(0)
        self._set_status(f"キュー: {os.path.basename(path)}", "working")
        self._on_play()

    def _benchmark_ollama(self) -> None:
        model = self.ollama_model_var.get().strip() or OLLAMA_MODEL_DEFAULT
        self._set_status(f"ベンチマーク中: {model} ...", "working")
        def _run():
            try:
                import time as _t
                t0 = _t.perf_counter()
                resp = ollama.chat(
                    model=model,
                    messages=[{"role": "user", "content": "日本語で「はい」とだけ答えてください。"}],
                )
                elapsed = _t.perf_counter() - t0
                content = (resp.message.content if hasattr(resp, "message")
                           else resp["message"]["content"]).strip()
                self.after(0, lambda: self._set_status(
                    f"{model}: {elapsed:.1f}秒 / 応答=「{content[:20]}」", "ok"))
            except Exception as exc:
                self.after(0, lambda: self._set_status(f"ベンチマーク失敗: {exc}", "error"))
        threading.Thread(target=_run, daemon=True).start()

    def _skip_patterns_text(self) -> str:
        pats = getattr(self, "_custom_skip_patterns", [])
        if not pats:
            return "パターンなし"
        return f"{len(pats)}件: " + " / ".join(pats[:3]) + ("..." if len(pats) > 3 else "")

    def _add_skip_pattern(self) -> None:
        pat = self._skip_pattern_entry.get().strip()
        if not pat:
            return
        try:
            re.compile(pat)
        except re.error as exc:
            self._set_status(f"無効な正規表現: {exc}", "error")
            return
        if pat not in self._custom_skip_patterns:
            self._custom_skip_patterns.append(pat)
            self._skip_pattern_entry.delete(0, "end")
            self._skip_pattern_label.configure(text=self._skip_patterns_text())
            self._save_settings()
            self._script_cache.clear()

    def _clear_skip_patterns(self) -> None:
        self._custom_skip_patterns.clear()
        self._skip_pattern_label.configure(text=self._skip_patterns_text())
        self._save_settings()
        self._script_cache.clear()

    def _apply_voicevox_port(self) -> None:
        global VOICEVOX_URL
        try:
            port = int(self.voicevox_port_var.get())
            VOICEVOX_URL = f"http://127.0.0.1:{port}"
            self._save_settings()
        except ValueError:
            self.voicevox_port_var.set(str(VOICEVOX_PORT))

    def _build_history_bar(self) -> None:
        self.history_frame = ctk.CTkFrame(self, corner_radius=6, height=32)
        self.history_frame.grid(row=4, column=0, sticky="ew", padx=12, pady=(4, 0))
        self._rebuild_history_bar()

    def _rebuild_history_bar(self) -> None:
        for w in self.history_frame.winfo_children():
            w.destroy()
        recent = self._recent_files[:5]
        if not recent:
            ctk.CTkLabel(
                self.history_frame, text="履歴なし",
                text_color="#666666", font=ctk.CTkFont(size=11)
            ).pack(side="left", padx=(10, 0), pady=4)
            return
        for path in recent:
            name = os.path.basename(path)
            label = name if len(name) <= 20 else name[:19] + "…"
            ctk.CTkButton(
                self.history_frame,
                text=f"📄 {label}",
                width=0,
                height=24,
                font=ctk.CTkFont(size=11),
                fg_color="#2B2B2B",
                hover_color="#3B3B3B",
                text_color="#AAAAAA",
                command=lambda p=path: self._on_import_history(p),
            ).pack(side="left", padx=(6, 0), pady=4)

    def _add_to_history(self, path: str) -> None:
        self._recent_files = [p for p in self._recent_files if p != path]
        self._recent_files.insert(0, path)
        self._recent_files = self._recent_files[:10]
        self._save_settings()
        self._rebuild_history_bar()

    def _on_import_history(self, path: str) -> None:
        if not os.path.exists(path):
            self._set_status(f"ファイルが見つかりません: {os.path.basename(path)}", "error")
            self._recent_files = [p for p in self._recent_files if p != path]
            self._save_settings()
            self._rebuild_history_bar()
            return
        try:
            text = _extract_text_from_file(
                path, ruby_to_kana=self.ruby_to_kana_var.get())
        except Exception as exc:
            self._set_status(f"読み込みエラー: {exc}", "error")
            return
        self.textbox.delete("1.0", "end")
        self.textbox.insert("1.0", text)
        self._script = []
        self._char_dict_dirty = True
        self._clear_all_highlights()
        self._update_chunks()
        self.time_slider.set(0)
        self._update_time_label(0)
        self._highlight_chunk(0)
        self._set_status(f"読み込み完了: {os.path.basename(path)}", "ok")
        self._add_to_history(path)
        self._restore_last_position()

    def _build_control_frame(self) -> None:
        frame = ctk.CTkFrame(self, corner_radius=8)
        frame.grid(row=3, column=0, sticky="ew", padx=12, pady=(6, 12))

        self.play_btn = ctk.CTkButton(
            frame, text="▶", width=44, command=self._on_play)
        self.play_btn.pack(side="left", padx=(10, 4), pady=8)

        self.pause_btn = ctk.CTkButton(
            frame, text="⏸", width=44,
            state="disabled", fg_color=COLOR_DISABLED, hover=False,
            command=self._on_pause_resume)
        self.pause_btn.pack(side="left", padx=(0, 4))

        self.stop_btn = ctk.CTkButton(
            frame, text="⏹", width=44,
            state="disabled", fg_color=COLOR_DISABLED, hover=False,
            command=self._on_stop)
        self.stop_btn.pack(side="left", padx=(0, 6))

        self.skip_btn_m100 = ctk.CTkButton(
            frame, text="⏮⏮", width=50, font=ctk.CTkFont(size=11),
            state="disabled", fg_color=COLOR_DISABLED, hover=False,
            command=lambda: self._skip_chunks(-100))
        self.skip_btn_m100.pack(side="left", padx=(6, 2))

        self.skip_btn_m10 = ctk.CTkButton(
            frame, text="⏮", width=44, font=ctk.CTkFont(size=13),
            state="disabled", fg_color=COLOR_DISABLED, hover=False,
            command=lambda: self._skip_chunks(-10))
        self.skip_btn_m10.pack(side="left", padx=(0, 2))

        self.skip_btn_p10 = ctk.CTkButton(
            frame, text="⏭", width=44, font=ctk.CTkFont(size=13),
            state="disabled", fg_color=COLOR_DISABLED, hover=False,
            command=lambda: self._skip_chunks(10))
        self.skip_btn_p10.pack(side="left", padx=(0, 2))

        self.skip_btn_p100 = ctk.CTkButton(
            frame, text="⏭⏭", width=50, font=ctk.CTkFont(size=11),
            state="disabled", fg_color=COLOR_DISABLED, hover=False,
            command=lambda: self._skip_chunks(100))
        self.skip_btn_p100.pack(side="left", padx=(0, 16))

        ctk.CTkLabel(frame, text="話速:").pack(side="left", padx=(0, 4))
        saved_speed = self._settings.get("playback_speed", SPEED_DEFAULT)
        self.speed_slider = ctk.CTkSlider(
            frame, from_=SPEED_MIN, to=SPEED_MAX, number_of_steps=25,
            width=150, command=self._on_speed_change)
        self.speed_slider.set(saved_speed)
        self.speed_slider.pack(side="left", padx=(0, 6))
        self.speed_label = ctk.CTkLabel(
            frame, text=f"{round(saved_speed, 1):.1f}x", width=50, anchor="w")
        self.speed_label.pack(side="left", padx=(0, 16))

        ctk.CTkLabel(frame, text="余白:").pack(side="left", padx=(0, 4))
        saved_padding = self._settings.get("silence_padding", SILENCE_PADDING_SEC)
        self.padding_slider = ctk.CTkSlider(
            frame, from_=0.0, to=1.0, number_of_steps=10,
            width=100, command=self._on_padding_change)
        self.padding_slider.set(saved_padding)
        self.padding_slider.pack(side="left", padx=(0, 4))
        self.padding_label = ctk.CTkLabel(
            frame, text=f"{saved_padding:.1f}s", width=40, anchor="w")
        self.padding_label.pack(side="left", padx=(0, 10))

        ctk.CTkLabel(frame, text="音量:").pack(side="left", padx=(10, 4))
        saved_volume = self._settings.get("playback_volume", VOLUME_DEFAULT)
        saved_boost  = self._settings.get("volume_boost", False)
        self.volume_boost_var = ctk.BooleanVar(value=saved_boost)
        slider_max = 2.0 if saved_boost else 1.0
        self.volume_slider = ctk.CTkSlider(
            frame, from_=0.0, to=slider_max, number_of_steps=20,
            width=100, command=self._on_volume_change)
        actual_volume = min(saved_volume, slider_max)
        self.volume_slider.set(actual_volume)
        self.volume_slider.pack(side="left", padx=(0, 4))
        self.volume_label = ctk.CTkLabel(
            frame, text=f"{round(actual_volume * 100):.0f}%", width=44, anchor="w")
        self.volume_label.pack(side="left", padx=(0, 4))
        self.volume_boost_btn = ctk.CTkButton(
            frame, text="", image=self._make_toggle_image(saved_boost),
            width=44, height=22, fg_color="transparent", hover=False,
            command=self._on_boost_btn_click)
        self.volume_boost_btn.pack(side="left", padx=(0, 10))

        self.normalize_var = ctk.BooleanVar(value=self._settings.get("normalize_volume", False))
        ctk.CTkCheckBox(frame, text="音量均一化", variable=self.normalize_var, width=90,
                        command=self._save_settings).pack(side="left", padx=(0, 10))

        ctk.CTkButton(
            frame, text="⏻ 終了", width=80,
            fg_color="#555555", hover_color="#333333",
            command=self._on_close,
        ).pack(side="right", padx=(0, 8))

        ctk.CTkButton(
            frame, text="💾 書出", width=72,
            command=self._export_script,
        ).pack(side="right", padx=(0, 4))

        ctk.CTkButton(
            frame, text="📋 ログ", width=72,
            command=self._open_log_window,
        ).pack(side="right", padx=(0, 4))

        ctk.CTkButton(
            frame, text="🔍 分割", width=72,
            command=self._show_chunk_preview,
        ).pack(side="right", padx=(0, 4))

        ctk.CTkButton(
            frame, text="🎵 WAV", width=72,
            command=self._save_wav,
        ).pack(side="right", padx=(0, 4))

        ctk.CTkButton(
            frame, text="📂 キュー", width=80,
            command=self._show_queue_window,
        ).pack(side="right", padx=(0, 4))

        self.status_label = ctk.CTkLabel(
            frame, text="エンジン起動中...", text_color=STATUS_COLORS["working"])
        self.status_label.pack(side="right", padx=14)

    # ─── チャンク管理 ─────────────────────────
    def _update_chunks(self) -> None:
        text = self.textbox.get("1.0", "end-1c")
        if self._script:
            self._chunks = [entry.get("text", "") for entry in self._script]
        else:
            self._chunks = self._split_text(text)
        self._chunk_tb_positions = self._compute_chunk_positions(text, self._chunks)
        n = len(self._chunks)
        if n > 1:
            self.time_slider.configure(from_=0, to=n - 1, number_of_steps=n - 1)
        else:
            self.time_slider.configure(from_=0, to=0, number_of_steps=1)
        self._update_time_label(int(self.time_slider.get()))
        self._update_char_count()

    def _compute_chunk_positions(
        self, text: str, chunks: list[str]
    ) -> list[tuple[str, str]]:
        positions: list[tuple[str, str]] = []
        search_start = 0
        for chunk in chunks:
            idx = text.find(chunk, search_start)
            if idx < 0:
                fallback = positions[-1] if positions else ("1.0", "1.0")
                positions.append(fallback)
                continue
            end_idx = idx + len(chunk)
            positions.append((
                self._offset_to_tb_index(text, idx),
                self._offset_to_tb_index(text, end_idx),
            ))
            search_start = end_idx
        return positions

    @staticmethod
    def _offset_to_tb_index(text: str, offset: int) -> str:
        before = text[:offset]
        line   = before.count("\n") + 1
        col    = len(before) - (before.rfind("\n") + 1)
        return f"{line}.{col}"

    def _highlight_chunk(self, chunk_idx: int) -> None:
        tb = self.textbox._textbox
        if (self._highlighted_chunk_idx >= 0
                and self._highlighted_chunk_idx < len(self._chunk_tb_positions)):
            old_start, old_end = self._chunk_tb_positions[self._highlighted_chunk_idx]
            if old_start != old_end:
                tb.tag_remove(HIGHLIGHT_TAG, old_start, old_end)
        if self._chunk_tb_positions and chunk_idx < len(self._chunk_tb_positions):
            start, end = self._chunk_tb_positions[chunk_idx]
            if start != end:
                tb.tag_add(HIGHLIGHT_TAG, start, end)
                tb.see(start)
        self._highlighted_chunk_idx = chunk_idx

    def _clear_all_highlights(self) -> None:
        self.textbox._textbox.tag_remove(HIGHLIGHT_TAG, "1.0", "end")
        self._highlighted_chunk_idx = -1

    def _update_time_label(self, idx: int) -> None:
        n = len(self._chunks)
        self.time_label.configure(text=f"{idx + 1 if n else 0} / {n}")

    # ─── イベントハンドラ ───────────────────
    def _on_text_changed(self, _=None) -> None:
        self._char_dict_dirty = True
        self._script = []
        self._script_cache.clear()
        self._update_char_count()

    def _update_char_count(self) -> None:
        text = self.textbox.get("1.0", "end-1c")
        n = len(text)
        label = f"{n:,} 文字"
        if n > 50_000:
            label += " ⚠ 大きなテキスト"
            self._char_count_label.configure(text=label, text_color="orange")
        else:
            self._char_count_label.configure(text=label, text_color="gray60")

    def _on_textbox_click(self, event=None) -> None:
        if not self._chunk_tb_positions:
            self._update_chunks()
        if not self._chunk_tb_positions:
            return
        try:
            cursor_pos = self.textbox._textbox.index("insert")
        except Exception:
            return

        # カーソル位置を含む最初のチャンクを探す
        clicked_idx = len(self._chunk_tb_positions) - 1
        for i, (start, end) in enumerate(self._chunk_tb_positions):
            try:
                if self.textbox._textbox.compare(cursor_pos, "<=", end):
                    clicked_idx = i
                    break
            except Exception:
                continue

        self._programmatic_slider_update = True
        if self._chunks:
            self.time_slider.set(clicked_idx)
        self._programmatic_slider_update = False
        self._update_time_label(clicked_idx)
        self._highlight_chunk(clicked_idx)

        if self._is_playing:
            self._seek_to(clicked_idx)

    def _on_clear(self) -> None:
        self.textbox.delete("1.0", "end")
        self._chunks = []
        self._chunk_tb_positions = []
        self._script = []
        self._char_dict_dirty = True
        self._clear_all_highlights()
        self.time_slider.configure(from_=0, to=1, number_of_steps=1)
        self.time_slider.set(0)
        self._update_time_label(0)

    def _get_clipboard_text(self) -> str:
        try:
            return self.clipboard_get()
        except Exception:
            return ""

    def _on_paste_and_reset(self) -> None:
        text = self._get_clipboard_text()
        if not text:
            self._set_status("クリップボードが空です", "error")
            return
        self.textbox.delete("1.0", "end")
        self.textbox.insert("1.0", text)
        self._script = []
        self._char_dict_dirty = True
        self._clear_all_highlights()
        self._update_chunks()
        self.time_slider.set(0)
        self._update_time_label(0)
        if self._chunks:
            self._highlight_chunk(0)
        self._set_status("ペースト完了", "ok")

    def _on_paste_and_play(self) -> None:
        self._on_paste_and_reset()
        if self._chunks:
            self._on_play()

    def _on_import(self) -> None:
        filetypes = [
            ("対応ファイル",       "*.txt *.md *.pdf *.docx *.epub"),
            ("テキストファイル",   "*.txt"),
            ("Markdown",          "*.md"),
            ("PDFファイル",        "*.pdf"),
            ("Word ファイル",      "*.docx"),
            ("EPUBファイル",       "*.epub"),
            ("すべてのファイル",   "*.*"),
        ]
        filepath = filedialog.askopenfilename(
            title="ファイルを選択してください", filetypes=filetypes)
        if not filepath:
            return
        try:
            text = _extract_text_from_file(
                filepath, ruby_to_kana=self.ruby_to_kana_var.get())
        except ImportError as exc:
            if DEBUG: print(f"[DEBUG] 読み込みエラー(ImportError): {exc}")
            self._set_status(str(exc), "error")
            return
        except Exception as exc:
            import traceback
            if DEBUG: print("[DEBUG] !!! 読み込みエラー !!!")
            traceback.print_exc()
            self._set_status(f"読み込みエラー: {exc}", "error")
            return

        self.textbox.delete("1.0", "end")
        self.textbox.insert("1.0", text)
        self._script = []
        self._char_dict_dirty = True
        self._clear_all_highlights()
        self._update_chunks()
        self.time_slider.set(0)
        self._update_time_label(0)
        self._highlight_chunk(0)
        self._set_status(f"読み込み完了: {os.path.basename(filepath)}", "ok")
        self._add_to_history(filepath)

    def _on_time_slider_change(self, value: float) -> None:
        if self._programmatic_slider_update:
            return
        idx = int(round(value))
        if self._chunks:
            idx = min(idx, len(self._chunks) - 1)
        self._update_time_label(idx)
        self._highlight_chunk(idx)
        if self._is_playing:
            self._seek_to(idx)

    def _on_speed_change(self, value: float) -> None:
        self.speed_label.configure(text=f"{round(value, 1):.1f}x")

    def _on_padding_change(self, value: float) -> None:
        self.padding_label.configure(text=f"{value:.1f}s")
        self._save_settings()

    def _on_volume_change(self, value: float) -> None:
        self.volume_label.configure(text=f"{round(value * 100):.0f}%")
        self._save_settings()

    def _on_volume_boost_toggle(self) -> None:
        boosted = self.volume_boost_var.get()
        new_max = 2.0 if boosted else 1.0
        current = self.volume_slider.get()
        if not boosted and current > 1.0:
            self.volume_slider.set(1.0)
            self.volume_label.configure(text="100%")
        self.volume_slider.configure(to=new_max, number_of_steps=20)
        self._save_settings()

    def _skip_chunks(self, delta: int) -> None:
        if not self._chunks:
            return
        current = max(self._highlighted_chunk_idx, 0)
        new_idx = max(0, min(len(self._chunks) - 1, current + delta))
        self._programmatic_slider_update = True
        self.time_slider.set(new_idx)
        self._programmatic_slider_update = False
        self._update_time_label(new_idx)
        self._highlight_chunk(new_idx)
        if self._is_playing:
            self._seek_to(new_idx)

    def _make_toggle_image(self, on: bool) -> ctk.CTkImage:
        from PIL import Image, ImageDraw
        w, h = 44, 22
        img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        draw.rounded_rectangle([0, 0, w - 1, h - 1],
                               radius=h // 2,
                               fill="#1F6AA5" if on else "#555555")
        d = h - 4
        cx = w - 2 - d if on else 2
        draw.ellipse([cx, 2, cx + d, h - 2], fill="white")
        return ctk.CTkImage(light_image=img, dark_image=img, size=(w, h))

    def _on_boost_btn_click(self) -> None:
        self.volume_boost_var.set(not self.volume_boost_var.get())
        self.volume_boost_btn.configure(
            image=self._make_toggle_image(self.volume_boost_var.get()))
        self._on_volume_boost_toggle()

    def _on_volume_up(self, _event=None) -> None:
        limit = 2.0 if self.volume_boost_var.get() else 1.0
        new_val = min(self.volume_slider.get() + 0.1, limit)
        self.volume_slider.set(new_val)
        self.volume_label.configure(text=f"{round(new_val * 100):.0f}%")
        self._save_settings()

    def _on_volume_down(self, _event=None) -> None:
        new_val = max(self.volume_slider.get() - 0.1, 0.0)
        self.volume_slider.set(new_val)
        self.volume_label.configure(text=f"{round(new_val * 100):.0f}%")
        self._save_settings()

    def _save_wav(self) -> None:
        if self._is_playing:
            self._set_status("再生中はWAV保存できません", "error")
            return
        if not self._chunks:
            self._set_status("テキストがありません", "error")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".wav",
            filetypes=[("WAV", "*.wav")],
            title="WAVファイルに保存",
        )
        if not path:
            return
        self._set_status("WAVを合成中... しばらくお待ちください", "working")
        self._set_controls_enabled(False)
        threading.Thread(target=self._save_wav_thread, args=(path,), daemon=True).start()

    def _save_wav_thread(self, output_path: str) -> None:
        import io as _io
        import wave as _wave
        try:
            chunks = self._script if self._script else [
                {"text": c, "style_id": 2} for c in self._chunks
            ]
            speed = round(self.speed_slider.get(), 1)
            volume = self.volume_slider.get()
            padding_sec = self.padding_slider.get()
            combined_chunks = []
            sample_rate_out = 24000
            for i, entry in enumerate(chunks):
                if not entry.get("text", "").strip():
                    continue
                self.after(0, lambda i=i, t=len(chunks): self._set_status(
                    f"WAV合成中... ({i+1}/{t})", "working"))
                try:
                    q_resp = self._http_session.post(
                        f"{VOICEVOX_URL}/audio_query",
                        params={"text": entry["text"], "speaker": int(entry.get("style_id", 2))},
                        timeout=30,
                    )
                    q_resp.raise_for_status()
                    aq = q_resp.json()
                    aq["postPhonemeLength"] = aq.get("postPhonemeLength", 0.0) + 0.15
                    s_resp = self._http_session.post(
                        f"{VOICEVOX_URL}/synthesis",
                        params={"speaker": int(entry.get("style_id", 2))},
                        json=aq, timeout=60,
                    )
                    s_resp.raise_for_status()
                    audio_data, sr = sf.read(_io.BytesIO(s_resp.content), dtype="float32")
                    sample_rate_out = sr
                    audio_data = _time_stretch(audio_data, sr, speed)
                    if getattr(self, "normalize_var", None) and self.normalize_var.get():
                        audio_data = _rms_normalize(audio_data)
                    if abs(volume - 1.0) > 0.01:
                        audio_data = np.clip(audio_data * volume, -1.0, 1.0)
                    silence = np.zeros(int(sr * padding_sec), dtype=np.float32)
                    combined_chunks.append(audio_data)
                    combined_chunks.append(silence)
                except Exception as exc:
                    if DEBUG: print(f"[DEBUG] WAV save: chunk {i} failed: {exc}")
            if not combined_chunks:
                self.after(0, lambda: self._set_status("合成できるチャンクがありません", "error"))
                return
            combined = np.concatenate(combined_chunks)
            sf.write(output_path, combined, sample_rate_out)
            self.after(0, lambda: self._set_status(
                f"WAV保存完了: {os.path.basename(output_path)}", "ok"))
        except Exception as exc:
            self.after(0, lambda: self._set_status(f"WAV保存エラー: {exc}", "error"))
        finally:
            self.after(0, lambda: self._set_controls_enabled(True))

    def _show_chunk_preview(self) -> None:
        if not self._chunks:
            self._update_chunks()
        if not self._chunks:
            self._set_status("チャンクがありません", "error")
            return
        win = ctk.CTkToplevel(self)
        win.title(f"チャンクプレビュー ({len(self._chunks)} 件)")
        win.geometry("600x480")
        win.grid_columnconfigure(0, weight=1)
        win.grid_rowconfigure(0, weight=1)
        scroll = ctk.CTkScrollableFrame(win)
        scroll.grid(row=0, column=0, sticky="nsew", padx=8, pady=(8, 4))
        scroll.grid_columnconfigure(1, weight=1)
        for i, chunk in enumerate(self._chunks):
            style_id = ""
            if self._script and i < len(self._script):
                style_id = str(self._script[i].get("style_id", ""))
            ctk.CTkLabel(scroll, text=f"{i+1}", width=45, anchor="e",
                         text_color="gray60", font=ctk.CTkFont(size=11)).grid(
                row=i, column=0, padx=(4, 6), pady=1, sticky="e")
            display = chunk[:80] + ("…" if len(chunk) > 80 else "")
            ctk.CTkLabel(scroll, text=display, anchor="w",
                         font=ctk.CTkFont(size=11)).grid(
                row=i, column=1, pady=1, sticky="w")
            if style_id:
                ctk.CTkLabel(scroll, text=style_id, width=50, anchor="e",
                             text_color="gray50", font=ctk.CTkFont(size=10)).grid(
                    row=i, column=2, padx=(4, 8), pady=1, sticky="e")
        ctk.CTkButton(win, text="閉じる", width=80,
                      command=win.destroy).grid(row=1, column=0, pady=(0, 8))

    def _export_script(self) -> None:
        import csv
        chunks = self._script if self._script else [
            {"text": c, "style_id": ""} for c in self._chunks
        ]
        if not chunks:
            self._set_status("エクスポートするデータがありません", "error")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("JSON", "*.json"), ("All", "*.*")],
            title="スクリプトを書き出す",
        )
        if not path:
            return
        try:
            if path.endswith(".json"):
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(chunks, f, ensure_ascii=False, indent=2)
            else:
                with open(path, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.writer(f)
                    writer.writerow(["index", "text", "style_id"])
                    for i, entry in enumerate(chunks):
                        writer.writerow([i + 1, entry.get("text", ""), entry.get("style_id", "")])
            self._set_status(f"書き出し完了: {os.path.basename(path)}", "ok")
        except OSError as exc:
            self._set_status(f"書き出しエラー: {exc}", "error")

    def _show_shortcuts(self) -> None:
        focus = self.focus_get()
        if focus and focus.winfo_class() in ("Entry", "Text"):
            return
        win = ctk.CTkToplevel(self)
        win.title("キーボードショートカット")
        win.geometry("360x280")
        win.resizable(False, False)
        shortcuts = [
            ("スペース",        "再生 / 一時停止"),
            ("Ctrl + ↑",       "音量 +10%"),
            ("Ctrl + ↓",       "音量 -10%"),
            ("?",              "このヘルプを表示"),
        ]
        frame = ctk.CTkFrame(win, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=20, pady=16)
        ctk.CTkLabel(frame, text="ショートカット一覧",
                     font=ctk.CTkFont(size=14, weight="bold")).pack(anchor="w", pady=(0, 10))
        for key, desc in shortcuts:
            row = ctk.CTkFrame(frame, fg_color="transparent")
            row.pack(fill="x", pady=3)
            ctk.CTkLabel(row, text=key, width=120, anchor="w",
                         font=ctk.CTkFont(family="Consolas", size=12)).pack(side="left")
            ctk.CTkLabel(row, text=desc, anchor="w").pack(side="left")
        ctk.CTkButton(win, text="閉じる", width=80,
                      command=win.destroy).pack(pady=(0, 12))
        win.grab_set()

    def _on_space_key(self) -> None:
        focus = self.focus_get()
        if focus and focus.winfo_class() in ("Entry", "Text"):
            return  # テキスト入力中はスペースを通常入力として扱う
        if self._is_playing:
            self._on_pause_resume()
        else:
            self._on_play()

    def _on_play(self) -> None:
        if self._is_llm_processing:
            if self._is_playing:
                return  # バックグラウンド配役中に再生中 → 既に動いているので何もしない
            # バックグラウンド配役中で停止中 → stop_eventを立てて中断し、再試行を促す
            self._stop_event.set()
            self._set_status("AI配役を中断中... もう一度▶を押してください", "working")
            return

        text = self.textbox.get("1.0", "end").strip()
        if not text:
            self._set_status("テキストを入力してください。", "error")
            return
        self._update_chunks()
        self._set_controls_enabled(False)
        self._stop_event.clear()

        start_index = min(int(round(self.time_slider.get())), max(0, len(self._chunks) - 1))
        current_tab = self.tabview.get()

        if DEBUG: print(f"\n[DEBUG] === Play Button Clicked ===")
        if DEBUG: print(f"[DEBUG] Current Tab Selected: '{current_tab}'")

        if current_tab == "AIディレクター (15役)":
            if DEBUG: print("[DEBUG] -> Triggering LLM Mode")
            threading.Thread(target=self._llm_and_play_thread, daemon=True).start()
        else:
            if DEBUG: print("[DEBUG] -> Triggering Rule-based Mode")
            self._script = []
            narrator_id = self._get_speaker_id(
                self.narrator_char_var.get(), self.narrator_style_var.get(), fallback=8)
            dialogue_id = self._get_speaker_id(
                self.dialogue_char_var.get(), self.dialogue_style_var.get(), fallback=2)
            chunks = self._split_text(text)
            self._script = [
                {"text": c, "style_id": _speaker_for_chunk(c, narrator_id, dialogue_id, self._valid_openers())}
                for c in chunks
            ]
            self._update_chunks()
            self._start_playback(start_index)

    # ─── ログウィンドウ ───────────────────────
    def _open_log_window(self) -> None:
        if self._log_window is not None and self._log_window.winfo_exists():
            self._log_window.focus()
            return
        win = ctk.CTkToplevel(self)
        win.title("デバッグログ")
        win.geometry("800x400")
        win.grid_columnconfigure(0, weight=1)
        win.grid_rowconfigure(0, weight=1)

        tb = ctk.CTkTextbox(win, font=ctk.CTkFont(family="Consolas", size=12),
                            wrap="word", state="disabled")
        tb.grid(row=0, column=0, sticky="nsew", padx=8, pady=(8, 4))

        ctk.CTkButton(win, text="クリア", width=80,
                      command=lambda: (tb.configure(state="normal"),
                                       tb.delete("1.0", "end"),
                                       tb.configure(state="disabled"))
                      ).grid(row=1, column=0, pady=(0, 8))

        self._log_window  = win
        self._log_textbox = tb
        win.protocol("WM_DELETE_WINDOW", self._close_log_window)

    def _close_log_window(self) -> None:
        self._log_textbox = None
        if self._log_window:
            self._log_window.destroy()
            self._log_window = None

    # ─── LLM前処理 (Stage1: 自動抽出, Stage2: バックグラウンド配役) ──
    def _llm_and_play_thread(self) -> None:
        self._is_llm_processing = True
        try:
            import traceback
            model = self.ollama_model_var.get().strip() or OLLAMA_MODEL_DEFAULT
            if DEBUG: print(f"\n[DEBUG] --- LLM Thread Start (Streaming Cast Mode) ---")

            if self._stop_event.is_set():
                return

            # ── Stage 1: 自動キャラクター抽出 (dirtyフラグが立っている場合のみ) ──
            stage1_profile_json: Optional[str] = None

            if self._char_dict_dirty:
                if DEBUG: print("[DEBUG] Dirty flag set → Stage 1: Auto Character Extraction")
                self.after(0, lambda: self._set_status("AI: 登場人物を自動抽出中...", "working"))

                full_text = self.textbox.get("1.0", "end").strip()
                if full_text:
                    max_chars = 8000
                    if len(full_text) > max_chars:
                        full_text = full_text[:max_chars] + "\n...[Text truncated]..."
                        if DEBUG: print(f"[DEBUG] Text truncated to {max_chars} chars.")
                    try:
                        resp = ollama.chat(
                            model=model,
                            messages=[
                                {"role": "system", "content": PROFILING_PROMPT},
                                {"role": "user",   "content": f"[Full Story Text]\n{full_text}"},
                            ],
                            format="json",
                            options={"num_ctx": 16384},
                        )
                        content = (resp.message.content if hasattr(resp, "message")
                                   else resp["message"]["content"]).strip()
                        if DEBUG: print(f"[DEBUG] Stage 1 raw:\n{content}\n" + "-" * 30)

                        if content:
                            if "```json" in content:
                                content = content.split("```json")[1].split("```")[0].strip()
                            elif "```" in content:
                                content = content.split("```")[1].split("```")[0].strip()
                            parsed = json.loads(content)
                            if isinstance(parsed, dict):
                                if "characters" in parsed and isinstance(parsed["characters"], list):
                                    extracted_list = parsed["characters"]
                                else:
                                    extracted_list = next(
                                        (v for v in parsed.values() if isinstance(v, list)), [])
                            elif isinstance(parsed, list):
                                extracted_list = parsed
                            else:
                                extracted_list = []

                            valid_items = [
                                item for item in extracted_list
                                if isinstance(item, dict) and "name" in item and "category" in item
                            ]
                            stage1_profile_json = json.dumps(
                                [{"name": str(i["name"]), "category": str(i["category"])}
                                 for i in valid_items],
                                ensure_ascii=False,
                            )
                            self.after(0, self._clear_all_rows)
                            for item in valid_items:
                                self.after(0, lambda n=str(item["name"]), c=str(item["category"]):
                                           self._add_char_row(n, c))
                            self._char_dict_dirty = False
                            if DEBUG: print(f"[DEBUG] Stage 1 complete: {len(valid_items)} characters.")
                        else:
                            if DEBUG: print("[DEBUG] Stage 1: empty response — using existing char dict.")
                    except Exception as exc:
                        if DEBUG: print(f"[DEBUG] Stage 1 failed: {exc} — using existing char dict.")

            if self._stop_event.is_set():
                self.after(0, lambda: self._set_status("LLM解析を中止しました", "stopped"))
                self.after(0, lambda: self._set_controls_enabled(True))
                return

            # ── プリフィル: ルールベースの仮speakerで全チャンクを初期化 ──
            narrator_id_llm = self._get_speaker_id(
                self.narrator_char_var.get(), self.narrator_style_var.get(), fallback=2)
            dialogue_id_llm = self._get_speaker_id(
                self.dialogue_char_var.get(), self.dialogue_style_var.get(), fallback=2)
            valid_openers = tuple(
                q.get()[0] for q in [self.quote1_var, self.quote2_var, self.quote3_var]
                if q.get() != "なし" and len(q.get()) == 2
            )
            if not valid_openers:
                valid_openers = ("「",)

            _cache_key = self._text_hash()
            if _cache_key in self._script_cache:
                self._script = [dict(e) for e in self._script_cache[_cache_key]]
                if DEBUG: print(f"[DEBUG] Script cache hit: {len(self._script)} chunks.")
                self.after(0, lambda: self._set_status("配役キャッシュを使用します", "working"))
            else:
                self._script = [
                    {
                        "text": c,
                        "style_id": _speaker_for_chunk(
                            c, narrator_id_llm, dialogue_id_llm, valid_openers),
                    }
                    for c in self._chunks
                ]
                if DEBUG: print(f"[DEBUG] Prefilled {len(self._script)} chunks with rule-based speakers.")

            # ── 即時再生開始 ──
            start_idx = min(
                int(round(self.time_slider.get())), max(0, len(self._chunks) - 1))
            self.after(0, lambda si=start_idx: self._start_playback(si))

            # ── Stage 2: バックグラウンドで逐次配役 ──
            character_profile_json = (
                stage1_profile_json if stage1_profile_json is not None
                else self._get_character_profile_json()
            )
            if DEBUG: print(f"[DEBUG] Stage 2 Background Cast. Profile: {character_profile_json[:80]}...")

            last_speaker_category = "主人公 女"
            total_chunks = len(self._chunks)
            cast_count = 0

            try:
                for i, chunk in enumerate(self._chunks):
                    if self._stop_event.is_set():
                        if DEBUG: print("[DEBUG] Stop event — background cast aborted.")
                        return

                    if not chunk.startswith(valid_openers):
                        continue

                    self.after(0, lambda idx=i, t=total_chunks: self._set_status(
                        f"AI: 配役中... ({idx + 1}/{t})", "working"))

                    start_ctx = max(0, i - 10)
                    end_ctx   = min(total_chunks, i + 11)
                    context_lines = []
                    for idx in range(start_ctx, end_ctx):
                        line = self._chunks[idx]
                        if idx == i:
                            context_lines.append(
                                f"===> [TARGET DIALOGUE TO CLASSIFY]: {line}")
                        else:
                            context_lines.append(line)
                    broad_context_str = "\n".join(context_lines)

                    system_msg = ATTRIBUTION_PROMPT.replace(
                        "{character_profile}", character_profile_json)
                    user_msg = (
                        f"[Broad Context (10 Chunks Before & After)]\n"
                        f"{broad_context_str}\n\n[Target Dialogue]\n{chunk}"
                    )

                    try:
                        resp = ollama.chat(
                            model=model,
                            messages=[
                                {"role": "system", "content": system_msg},
                                {"role": "user",   "content": user_msg},
                            ],
                            format={
                                "type": "object",
                                "properties": {"category": {"type": "string"}},
                                "required": ["category"],
                            },
                            options={"num_ctx": 8192},
                        )
                        if self._stop_event.is_set():
                            return
                        cat_content = (resp.message.content if hasattr(resp, "message")
                                       else resp["message"]["content"])
                        parsed_cat = json.loads(cat_content)
                        category   = str(parsed_cat.get(
                            "category", last_speaker_category)).strip()
                        if category == "ナレーション" or category not in self.archetype_vars:
                            category = last_speaker_category
                    except Exception as exc:
                        if DEBUG: print(f"[DEBUG] Error classifying chunk {i}: {exc}")
                        category = last_speaker_category

                    last_speaker_category = category

                    if category in self.archetype_vars:
                        c_name   = self.archetype_vars[category]["char"].get()
                        s_name   = self.archetype_vars[category]["style"].get()
                        style_id = self._get_speaker_id(c_name, s_name, fallback=2)
                    else:
                        c_name   = self.narrator_char_var.get()
                        s_name   = self.narrator_style_var.get()
                        style_id = self._get_speaker_id(c_name, s_name, fallback=2)

                    self._script[i]["style_id"] = style_id
                    self._script[i]["category"] = category
                    cast_count += 1
                    if DEBUG: print(f"[DEBUG] Chunk {i}: Cat='{category}', ID={style_id}")

            except Exception:
                if DEBUG: print("\n[DEBUG] !!! STAGE 2 CRITICAL ERROR !!!")
                if DEBUG: traceback.print_exc()
                # 再生はすでに始まっているので止めない。ステータスだけ更新。
                self.after(0, lambda: self._set_status("AI配役でエラーが発生しました", "error"))
                return

            self._script_cache[_cache_key] = [dict(e) for e in self._script]
            self.after(0, lambda n=cast_count: self._set_status(
                f"AI配役完了 ({n}チャンク)", "ok"))
            if DEBUG: print(f"[DEBUG] Background cast complete: {cast_count} dialogues classified.")

        except Exception:
            import traceback
            if DEBUG: print("\n[DEBUG] !!! LLM THREAD CRITICAL ERROR !!!")
            if DEBUG: traceback.print_exc()
            self._script = []
            self.after(0, lambda: self._set_status("LLM解析に失敗しました", "error"))
            if not self._is_playing:
                self.after(0, lambda: self._set_controls_enabled(True))
        finally:
            self._is_llm_processing = False

    def _on_llm_done(self) -> None:
        self._update_chunks()
        start_index = 0
        if self._chunks:
            raw_idx = int(round(self.time_slider.get()))
            start_index = min(raw_idx, len(self._chunks) - 1)
        self._start_playback(start_index)

    def _start_playback(self, start_index: int) -> None:
        text = self.textbox.get("1.0", "end").strip()
        narrator_id = self._get_speaker_id(
            self.narrator_char_var.get(), self.narrator_style_var.get(), fallback=8)
        dialogue_id = self._get_speaker_id(
            self.dialogue_char_var.get(), self.dialogue_style_var.get(), fallback=2)

        self._stop_event.clear()
        self._pause_event.clear()
        self._is_paused  = False
        self._is_playing = True

        self._cleanup_temp_files()
        self._audio_queue    = queue.Queue()
        self._slot_semaphore = threading.Semaphore(4)

        self._play_generation += 1
        gen = self._play_generation

        self._set_controls_enabled(False)
        n = len(self._chunks)
        self._set_status(f"処理を開始します... ({start_index + 1}/{n})", "working")

        self._producer_thread = threading.Thread(
            target=self._producer,
            args=(text, narrator_id, dialogue_id, start_index, gen),
            daemon=True,
        )
        self._consumer_thread = threading.Thread(
            target=self._consumer,
            args=(gen,),
            daemon=True,
        )
        self._producer_thread.start()
        self._consumer_thread.start()

    def _on_pause_resume(self) -> None:
        if not self._is_paused:
            self._is_paused = True
            self._pause_event.set()
            sd.stop()
            self.pause_btn.configure(text="▶")
            self._set_status("一時停止中", "paused")
        else:
            self._is_paused = False
            self._pause_event.clear()
            self.pause_btn.configure(text="⏸")
            self._set_status("再開します...", "working")

    def _on_stop(self) -> None:
        self._save_last_position()
        self._stop_event.set()
        self._pause_event.clear()
        self._is_paused = False
        sd.stop()
        if self._audio_queue is not None:
            try:
                self._audio_queue.put_nowait(None)
            except Exception:
                pass
        self._set_status("停止しました", "stopped")

    def _seek_to(self, idx: int) -> None:
        self._play_generation += 1
        seek_gen = self._play_generation
        self._seek_pending_gen = seek_gen

        self._stop_event.set()
        sd.stop()

        if self._audio_queue is not None:
            try:
                self._audio_queue.put_nowait(None)
            except Exception:
                pass
            while True:
                try:
                    self._audio_queue.get_nowait()
                except queue.Empty:
                    break

        threading.Thread(
            target=self._wait_and_restart,
            args=(idx, seek_gen),
            daemon=True,
        ).start()

    def _wait_and_restart(self, idx: int, seek_gen: int) -> None:
        time.sleep(0.1)
        if self._producer_thread and self._producer_thread.is_alive():
            self._producer_thread.join(timeout=2.0)
        if self._consumer_thread and self._consumer_thread.is_alive():
            self._consumer_thread.join(timeout=2.0)

        if seek_gen != self._seek_pending_gen:
            return
        if not self._is_playing:
            return
        self.after(0, lambda: self._apply_seek(idx, seek_gen))

    def _apply_seek(self, idx: int, seek_gen: int) -> None:
        if seek_gen != self._seek_pending_gen:
            return
        if not self._is_playing:
            return
        self._update_chunks()
        idx = min(idx, max(0, len(self._chunks) - 1))
        self._start_playback(idx)

    # ─── プロデューサースレッド (合成) ─────────
    def _producer(
        self,
        text: str,
        narrator_id: int,
        dialogue_id: int,
        start_index: int,
        gen: int,
    ) -> None:
        assert self._audio_queue is not None
        try:
            # エンジン/キャプションは開始時に一度だけスナップショット（スレッド安全・セッション内固定）
            _engine = self.engine_var.get() if hasattr(self, "engine_var") else "voicevox"
            _caption_map = ({c: v.get() for c, v in self.caption_vars.items()}
                            if hasattr(self, "caption_vars") else {})
            _narr_caption = (self.narrator_caption_var.get()
                             if hasattr(self, "narrator_caption_var") else "")
            _use_ref = (self.use_ref_var.get() if hasattr(self, "use_ref_var")
                        else self._settings.get("irodori_use_ref", True))
            _seeds = dict(self._caption_seeds)
            _global_engine = _engine  # "voicevox" / "irodori" / "mixed"
            _cat_engines = dict(self._category_engines)
            # Irodori を1チャンクでも使うならバンドルを遅延起動（mixed 含む）
            _needs_irodori = (_global_engine == "irodori"
                              or (_global_engine == "mixed"
                                  and any(v == "irodori" for v in _cat_engines.values())))
            # Irodori 起動はUIを固めぬよう producer スレッドで実行
            if _needs_irodori:
                try:
                    ok = self._irodori.ensure_running(
                        on_status=lambda m, k="working": self.after(
                            0, lambda mm=m, kk=k: self._set_status(mm, kk)))
                except irodori_engine.IrodoriLaunchError as exc:
                    self.after(0, lambda e=str(exc): self._set_status(e, "error"))
                    self._stop_event.set()
                    self._audio_queue.put(None)
                    return
                if not ok or self._stop_event.is_set() or gen != self._play_generation:
                    if not ok:
                        self.after(0, lambda: self._set_status("Irodori 起動に失敗しました", "error"))
                        self._stop_event.set()
                    self._audio_queue.put(None)
                    return

            if self._script:
                total = len(self._script)

                def _iter_script():
                    for ci in range(total):
                        entry = self._script[ci]
                        tts_text = entry.get("text", "")
                        try:
                            spk = int(entry.get("style_id", narrator_id))
                        except (ValueError, TypeError):
                            spk = narrator_id
                        yield ci, tts_text, spk

                _chunks_iter = _iter_script()
            else:
                raw = self._split_text(text)

                def _iter_rule():
                    for ci, c in enumerate(raw):
                        spk = _speaker_for_chunk(
                            c, narrator_id, dialogue_id, self._valid_openers())
                        yield ci, c, spk

                _chunks_iter = _iter_rule()
                total = len(raw)

            for chunk_idx, chunk, speaker_id in _chunks_iter:
                index = chunk_idx + 1

                if chunk_idx < start_index:
                    continue

                if self._stop_event.is_set() or gen != self._play_generation:
                    break

                while self._pause_event.is_set():
                    if self._stop_event.is_set() or gen != self._play_generation:
                        break
                    time.sleep(0.05)
                if self._stop_event.is_set() or gen != self._play_generation:
                    break

                slot_path = TEMP_PATHS[(chunk_idx - start_index) % TEMP_SLOT_COUNT]
                slot_acquired = False
                while True:
                    if self._stop_event.is_set() or gen != self._play_generation:
                        break
                    if self._slot_semaphore.acquire(timeout=0.1):
                        slot_acquired = True
                        break
                if not slot_acquired:
                    break

                if gen != self._play_generation:
                    self._slot_semaphore.release()
                    break

                if not chunk.endswith(("。", "！", "？", "!", "?")):
                    chunk += "。"

                try:
                    # チャンクのカテゴリ → 使用エンジン（global=mixed のときカテゴリ別）
                    if self._script and chunk_idx < len(self._script):
                        cat = self._script[chunk_idx].get("category", "ナレーション")
                    else:
                        cat = "ナレーション"
                    _chunk_engine = irodori_engine.engine_for(cat, _global_engine, _cat_engines)
                    if _chunk_engine == "irodori":
                        caption = irodori_engine.resolve_caption(
                            cat, _caption_map, _narr_caption)
                        seed = irodori_engine.voice_seed_for(cat, _seeds)
                        self.after(0, lambda i=index, t=total: self._set_status(
                            f"Irodori 合成中... ({i}/{t})", "working"))
                        # カテゴリ毎seed＋use_ref（声固定）で同カテゴリの声を一貫させる
                        wav_bytes = irodori_engine.synthesize_irodori(
                            self._http_session, self._irodori.base_url, chunk, caption,
                            seed=seed, use_ref=_use_ref)
                    else:
                        self.after(0, lambda i=index, t=total: self._set_status(
                            f"音声クエリを送信中... ({i}/{t})", "working"))
                        query_resp = self._http_session.post(
                            f"{VOICEVOX_URL}/audio_query",
                            params={"text": chunk, "speaker": speaker_id},
                            timeout=30,
                        )
                        query_resp.raise_for_status()
                        audio_query = query_resp.json()
                        del query_resp
                        audio_query["postPhonemeLength"] = (
                            audio_query.get("postPhonemeLength", 0.0) + 0.15
                        )

                        self.after(0, lambda i=index, t=total: self._set_status(
                            f"WAVデータを合成中... ({i}/{t})", "working"))
                        synthesis_resp = self._http_session.post(
                            f"{VOICEVOX_URL}/synthesis",
                            params={"speaker": speaker_id},
                            json=audio_query,
                            timeout=60,
                        )
                        synthesis_resp.raise_for_status()
                        del audio_query
                        wav_bytes = synthesis_resp.content
                        del synthesis_resp

                    if gen != self._play_generation:
                        self._slot_semaphore.release()
                        break

                    write_ok = False
                    for _attempt in range(5):
                        try:
                            with open(slot_path, "wb") as f:
                                f.write(wav_bytes)
                            write_ok = True
                            break
                        except PermissionError:
                            time.sleep(0.05)
                    del wav_bytes
                    if not write_ok:
                        raise PermissionError(
                            f"一時ファイルへの書き込みが5回失敗しました: "
                            f"{os.path.basename(slot_path)}"
                        )

                    self._audio_queue.put((slot_path, index, total, chunk_idx))

                except (requests.exceptions.ConnectionError,
                        requests.exceptions.Timeout,
                        requests.exceptions.HTTPError):
                    self._slot_semaphore.release()
                    raise  # 致命的エラーは外側 except に委ねる
                except Exception as exc:
                    self._slot_semaphore.release()
                    msg = str(exc)[:60]
                    self.after(0, lambda m=msg: self._set_status(
                        f"チャンクをスキップ: {m}", "working"))
                    self._audio_queue.put(("SKIP", index, total, chunk_idx))

        except requests.exceptions.ConnectionError:
            self.after(0, lambda: self._set_status(
                "接続エラー: VOICEVOX engine.exe を起動してください（ポート50021）", "error"))
            self._stop_event.set()
        except requests.exceptions.Timeout:
            self.after(0, lambda: self._set_status(
                "タイムアウト: VOICEVOXの応答が遅延しています", "error"))
            self._stop_event.set()
        except requests.exceptions.HTTPError as exc:
            code = exc.response.status_code if exc.response is not None else "?"
            self.after(0, lambda: self._set_status(f"APIエラー: HTTP {code}", "error"))
            self._stop_event.set()
        except Exception as exc:
            msg = str(exc)
            self.after(0, lambda: self._set_status(f"予期しないエラー: {msg}", "error"))
            self._stop_event.set()
        finally:
            self._audio_queue.put(None)

    # ─── コンシューマースレッド (再生) ─────────
    def _consumer(self, gen: int) -> None:
        assert self._audio_queue is not None

        try:
            while True:
                while self._pause_event.is_set():
                    if self._stop_event.is_set():
                        break
                    time.sleep(0.05)

                if self._stop_event.is_set():
                    break

                item = self._audio_queue.get()
                if item is None:
                    break

                if isinstance(item, tuple) and item[0] == "SKIP":
                    _, skip_i, skip_t, skip_ci = item
                    self.after(0, lambda ci=skip_ci: self._on_consumer_progress(ci))
                    continue

                slot_path, index, total, chunk_idx = item

                self.after(0, lambda ci=chunk_idx: self._on_consumer_progress(ci))
                self.after(0, lambda i=index, t=total: self._set_status(
                    f"再生中... ({i}/{t})", "working"))

                audio_data: Optional[np.ndarray] = None
                sample_rate: Optional[int] = None
                for _attempt in range(5):
                    try:
                        audio_data, sample_rate = sf.read(slot_path, dtype="float32")
                        break
                    except FileNotFoundError:
                        break
                    except (PermissionError, OSError):
                        time.sleep(0.05)

                self._slot_semaphore.release()

                if audio_data is None:
                    break

                # 再生レートを統一（Irodoriの48kHz等は一部環境で sd.play が詰まるため24kHzへ）。
                # VOICEVOXの24kHzはそのまま。speech/感情表現は24kHzで実質劣化なし。
                if sample_rate != PLAYBACK_SR:
                    audio_data = _resample_linear(audio_data, sample_rate, PLAYBACK_SR)
                    sample_rate = PLAYBACK_SR

                speed = round(self.speed_slider.get(), 1)
                audio_processed = _time_stretch(audio_data, sample_rate, speed)
                del audio_data
                if getattr(self, "normalize_var", None) and self.normalize_var.get():
                    audio_processed = _rms_normalize(audio_processed)

                pre_len     = int(sample_rate * PRE_PADDING_SEC)
                silence_len = int(sample_rate * self.padding_slider.get())
                if audio_processed.ndim == 1:
                    pre_silence = np.zeros(pre_len,     dtype=np.float32)
                    silence     = np.zeros(silence_len, dtype=np.float32)
                else:
                    pre_silence = np.zeros(
                        (pre_len,     audio_processed.shape[1]), dtype=np.float32)
                    silence     = np.zeros(
                        (silence_len, audio_processed.shape[1]), dtype=np.float32)
                padded = np.concatenate([pre_silence, audio_processed, silence])
                del pre_silence, audio_processed, silence

                volume = self.volume_slider.get()
                if abs(volume - 1.0) > 0.01:
                    padded = np.clip(padded * volume, -1.0, 1.0)

                sd.play(padded, samplerate=sample_rate)
                sd.wait()
                del padded

                if self._stop_event.is_set():
                    break

            if gen == self._play_generation and not self._stop_event.is_set():
                self.after(0, lambda: self._set_status("完了", "ok"))

        except Exception as exc:
            msg = str(exc)
            if gen == self._play_generation:
                self.after(0, lambda: self._set_status(f"再生エラー: {msg}", "error"))
        finally:
            if gen == self._play_generation:
                self._is_playing = False
                self._cleanup_temp_files()
                self._is_paused = False
                self._pause_event.clear()
                self.after(0, lambda: self.pause_btn.configure(text="⏸"))
                self.after(0, lambda: self._set_controls_enabled(True))
                self.after(0, self._clear_all_highlights)

    def _on_consumer_progress(self, chunk_idx: int) -> None:
        self._programmatic_slider_update = True
        if self._chunks:
            self.time_slider.set(chunk_idx)
        self._update_time_label(chunk_idx)
        self._programmatic_slider_update = False
        self._highlight_chunk(chunk_idx)

    # ─── ヘルパー ────────────────────────────
    def _set_status(self, message: str, level: str = "idle") -> None:
        color = STATUS_COLORS.get(level, STATUS_COLORS["idle"])
        self.status_label.configure(text=message, text_color=color)

    def _set_controls_enabled(self, enabled: bool) -> None:
        idle_state = "normal" if enabled else "disabled"
        for widget in (
            self.play_btn, self.paste_btn, self.paste_play_btn,
            self.clear_btn, self.import_btn,
            self.narrator_char_menu, self.narrator_style_menu,
            self.dialogue_char_menu, self.dialogue_style_menu,
            *self._archetype_menus,
        ):
            widget.configure(state=idle_state)
        self.speed_slider.configure(state="normal")
        self.volume_slider.configure(state="normal")
        self.time_slider.configure(state="normal")
        if enabled:
            for btn in (self.pause_btn, self.stop_btn):
                btn.configure(state="disabled", fg_color=COLOR_DISABLED, hover=False)
            for btn in (self.skip_btn_m100, self.skip_btn_m10,
                        self.skip_btn_p10,  self.skip_btn_p100):
                btn.configure(state="disabled", fg_color=COLOR_DISABLED, hover=False)
        else:
            self.pause_btn.configure(
                state="normal", fg_color=COLOR_PAUSE_ACTIVE, hover=True)
            self.stop_btn.configure(
                state="normal", fg_color=COLOR_STOP_ACTIVE, hover=True)
            for btn in (self.skip_btn_m100, self.skip_btn_m10,
                        self.skip_btn_p10,  self.skip_btn_p100):
                btn.configure(state="normal", fg_color=COLOR_PAUSE_ACTIVE, hover=True)

    def _cleanup_temp_files(self) -> None:
        for path in TEMP_PATHS:
            if os.path.exists(path):
                try:
                    os.unlink(path)
                except OSError:
                    pass

    def destroy(self) -> None:
        self._stop_event.set()
        self._pause_event.clear()
        sd.stop()
        if self._audio_queue is not None:
            try:
                self._audio_queue.put_nowait(None)
            except Exception:
                pass
        self._cleanup_temp_files()
        self._terminate_engine()
        self._http_session.close()
        super().destroy()


# ─────────────────────────────────────────
# エントリーポイント
# ─────────────────────────────────────────
if __name__ == "__main__":
    ctk.set_appearance_mode("dark")
    ctk.set_default_color_theme("blue")
    app = VoicevoxTTSApp()
    app.mainloop()
